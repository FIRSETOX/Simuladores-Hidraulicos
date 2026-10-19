"""
MOTOR DEL SIMULADOR HIDRÁULICO EN ESTADO ESTABLE (sin interfaz gráfica).

Reproduce las fórmulas de la plantilla MLO-FR-05 "PLANTILLA SIMULADOR HIDRAULICO.xlsm":
cálculo por tramo, resumen, interpolación, curva del sistema, Buscar objetivo,
actualización de espesores, importación de plantillas y base de datos de sistemas.
Lo usa la aplicación web (app.py) y se puede usar desde cualquier script de Python.
"""
import csv
import getpass
import io
import json
import math
import os
import re
import sqlite3
from dataclasses import dataclass, field, asdict, replace
from datetime import datetime










# --------------------------------------------------------------------------
# Constantes (las mismas que usa la hoja de cálculo)
# --------------------------------------------------------------------------
PSI_POR_M_KGM3 = 0.00142245          # J = altura * densidad * 0.00142245
KBD_A_M3S = 35000 / 220.5 / 24 / 3600  # B7 = B6*35000/220.5/24/3600
KBD_A_BPH = 1000 / 24                # B8 = B6*1000/24
BPH_A_KBD = 0.024                    # Resumen E15: inyección (bph) * 0.024
IN_A_M = 0.0254
IN_A_MM = 25.4
FACTOR_DISENO = 0.72                 # MAOP = 0.72*2*SMYS*t/D
RE_LAMINAR = 2200

INYECCIONES = ("Sin inyección", "INY 1", "INY 2")

# Tabla REF de la plantilla (Nombre, Densidad kg/m3, Visc cSt, T1 °F, Dens2, Visc2, T2)
PRODUCTOS_DEFECTO = [
    ["GM", 740, 0.6, 86, None, None, None],
    ["B2E", 870, 4.6, 86, None, None, None],
    ["JET", 821, 1.64, 86, None, None, None],
    ["VO", 780, 1.2, 86, None, None, None],
    ["KDHT", 820, 1.64, 86, None, None, None],
    ["CRUDO CASTILLA NORTE", 957.095801578639, 722, 86, 950.696303347836, 358.8, 104],
    ["MEZCLA APIAY", 952.685423160279, 650, 86, 906.308250949371, 450, 104],
    ["Andalucia", 871.4, 7.88, 104, None, 6.19, 122],
    ["Vasconia", 907.9, 22.02, 104, None, None, None],
    ["Castilla", 980, 2297.88, 104, None, None, None],
    ["Mezcla Castilla-Vasconia", 941.15, 174.67, 104, None, None, None],
    ["Caño Limon", 876, 12.05, 104, None, None, None],
    ["Nafta", 704, 0.504, 104, None, None, None],
    ["GLP", 595, 0.27915966386554625, 85, None, None, None],
    ["Floreña", 794, 1.72, 104, None, None, None],
    ["Southblend", 877, 9.82, 104, None, None, None],
    ["Combustoleo", 1020, 350, 86, None, None, None],
    ["Agua", 1000, 1, 86, None, None, None],
]


# --------------------------------------------------------------------------
# Estructuras de datos
# --------------------------------------------------------------------------
@dataclass
class PuntoPerfil:
    punto: str          # D  Punto/ubicación
    km: float           # E  dist. del reg. [km]
    altura: float       # F  altura [m]
    espesor: float      # G  esp. [in]
    diam_ext: float     # H  diám. ext. [in]
    smys: float         # I  SMYS [psi]
    rating: float = 0.0  # U  Rating de válvula [psi] (0 = usar MAOP por espesor)


@dataclass
class Tramo:
    nombre: str
    puntos: list = field(default_factory=list)   # list[PuntoPerfil]
    rugosidad_in: float = 0.0018                 # B18
    eficiencia: float = 100.0                    # B20
    formato_antiguo: bool = False                # plantillas antiguas: fórmula de fricción de esa versión
    fila_inicio: int = 3                         # fila de Excel del primer punto (para referencia)
    dra_siempre: bool = False                    # formato con Barlow: el %DR se aplica también en laminar
    factor_q: float = 35000 / 220.5 / 24 / 3600  # m3/s por KBDO (la plantilla Barlow usa 0.00004416 por bph)


@dataclass
class Estacion:
    nombre: str
    presion_recibo: float = 0.0      # Resumen col C
    paso_directo: bool = False       # True -> recibo = descarga del tramo siguiente (C17 = D17)
    dr: float = 0.0                  # %DR del tramo que inicia en esta estación (fila 10)
    fr: float = 1 / 1.1              # Factor de rotura (fila 11)
    inyeccion: str = "Sin inyección"  # Resumen col J
    delta_bph: float = 0.0           # inyección (+) / entrega al paso (−) en la estación, bph
    clientes: str = ""               # entregas a clientes (solo para el esquema), p.ej. "BIOMAX, PUMA"


@dataclass
class Proyecto:
    sistema: str = ""
    fecha: str = ""
    responsable: str = ""
    flujo_kbd: float = 86.4
    producto: str = "GM"
    densidad: float = 740.0
    viscosidad: float = 0.6
    iny1_bph: float = 0.0
    iny2_bph: float = 0.0
    estaciones: list = field(default_factory=list)  # n+1 estaciones
    tramos: list = field(default_factory=list)      # n tramos
    productos: list = field(default_factory=lambda: [list(p) for p in PRODUCTOS_DEFECTO])
    km_interpolar: list = field(default_factory=lambda: [5, 20, 90, 100])
    perfiles_externos: list = field(default_factory=list)   # perfiles MOP / transitorios cargados
    margen_descuelgue_psi: float = 0.0   # presión interna mínima exigida contra descuelgue (0 = cabeza = altura)
    ramales: list = field(default_factory=list)   # list[Ramal]: derivaciones del flujo (red de tramos)

    # ---------------- serialización ----------------
    def to_dict(self):
        return asdict(self)

    @staticmethod
    def from_dict(d):
        from dataclasses import fields as _campos
        validos = {f.name for f in _campos(Proyecto)} - {"estaciones", "tramos", "ramales"}
        p = Proyecto(**{k: v for k, v in d.items() if k in validos})
        p.estaciones = estaciones_desde(d.get("estaciones", []))
        p.tramos = tramos_desde(d.get("tramos", []))
        p.ramales = [Ramal.from_dict(r) for r in d.get("ramales", []) or []]
        return p


def estaciones_desde(lista):
    from dataclasses import fields as _campos
    campos_est = {f.name for f in _campos(Estacion)}
    return [e if isinstance(e, Estacion) else Estacion(**{k: v for k, v in e.items() if k in campos_est})
            for e in lista]


def tramos_desde(lista):
    salida = []
    for t in lista:
        if isinstance(t, Tramo):
            salida.append(t)
            continue
        tr = Tramo(nombre=t["nombre"], rugosidad_in=t.get("rugosidad_in", 0.0018),
                   eficiencia=t.get("eficiencia", 100.0),
                   formato_antiguo=t.get("formato_antiguo", t.get("dra_en_laminar", False)),
                   fila_inicio=t.get("fila_inicio", 3), dra_siempre=t.get("dra_siempre", False),
                   factor_q=t.get("factor_q", KBD_A_M3S))
        tr.puntos = [p if isinstance(p, PuntoPerfil) else PuntoPerfil(**p) for p in t.get("puntos", [])]
        salida.append(tr)
    return salida


MODOS_RAMAL = ("Porcentaje", "Hidráulico")


@dataclass
class Ramal:
    """Derivación del flujo: cadena de tramos que sale de una estación (nodo) de la línea principal o de
    otro ramal. estaciones[0] es el nodo (su %DR y FR aplican al primer tramo del ramal).

    modo "Porcentaje": el ramal se lleva pct % del flujo que llega al nodo (válvula de control); la
        presión en el nodo es la mayor que exijan los ramales y la otra rama estrangula la diferencia.
    modo "Hidráulico": el reparto lo da la hidráulica (misma presión en el nodo para ambas ramas,
        con las presiones de recibo de cada destino)."""
    nombre: str
    padre: str = ""                   # "" = línea principal; si no, nombre del ramal padre
    nodo: int = 0                     # índice de la estación del padre donde se deriva
    modo: str = "Porcentaje"
    pct: float = 50.0
    estaciones: list = field(default_factory=list)
    tramos: list = field(default_factory=list)

    @staticmethod
    def from_dict(d):
        r = Ramal(nombre=d.get("nombre", "Ramal"), padre=d.get("padre", "") or "", nodo=int(d.get("nodo", 0)),
                  modo=d.get("modo", "Porcentaje"), pct=float(d.get("pct", 50.0)))
        r.estaciones = estaciones_desde(d.get("estaciones", []))
        r.tramos = tramos_desde(d.get("tramos", []))
        return r


# --------------------------------------------------------------------------
# Resultados
# --------------------------------------------------------------------------
@dataclass
class ResultadoTramo:
    nombre: str
    flujo_kbd: float
    caudal_m3s: float
    presion_recibo: float
    dr: float
    fr: float
    J: list = field(default_factory=list)   # presión geodésica
    K: list = field(default_factory=list)   # velocidad
    L: list = field(default_factory=list)   # Reynolds
    M: list = field(default_factory=list)   # pérdidas psi/km
    N: list = field(default_factory=list)   # presión int + geod
    O: list = field(default_factory=list)   # presión int
    P: list = field(default_factory=list)   # MAOP
    Q: list = field(default_factory=list)   # MAOP + geod
    R: list = field(default_factory=list)   # descuelgue
    S: list = field(default_factory=list)   # rotura de tubo
    T: list = field(default_factory=list)   # lleno de línea acumulado (bls)

    @property
    def presion_descarga(self):
        return self.O[0] if self.O else 0.0

    @property
    def roturas(self):
        return sum(1 for s in self.S if s)

    @property
    def descuelgues(self):
        return sum(1 for r in self.R if r)

    @property
    def lleno_kbls(self):
        return (self.T[-1] if self.T else 0.0) / 1000

    @property
    def presion_remanente(self):
        """Resumen col I: MIN de (N-J) positivos."""
        pos = [o for o in self.O if o > 0]
        return min(pos) if pos else None

    @property
    def re_min(self):
        vals = [l for l in self.L if l > 0]
        return min(vals) if vals else 0.0

    @property
    def regimen(self):
        re_ = self.re_min
        if re_ < RE_LAMINAR:
            return "Laminar"
        if re_ < 4500:
            return "Transicional (revisar %DR)"
        return "Turbulento"


# --------------------------------------------------------------------------
# Cálculo
# --------------------------------------------------------------------------
def factor_friccion_churchill(re, rug_mm, d_mm):
    """Ecuación de Churchill tal como está en la columna M."""
    if re < 1.0:                     # caudal prácticamente nulo: flujo laminar (Darcy 64/Re = 8·(8/Re))
        return 64 / re
    a = (2.457 * math.log(1 / ((7 / re) ** 0.9 + 0.27 * rug_mm / d_mm))) ** 16
    b = (37530 / re) ** 16
    return 8 * ((8 / re) ** 12 + 1 / (a + b) ** 1.5) ** (1 / 12)


def calcular_tramo(tramo: Tramo, flujo_kbd, densidad, visc_cst, dr, fr, presion_recibo):
    n = len(tramo.puntos)
    q_si = flujo_kbd * tramo.factor_q
    mu = visc_cst * densidad / 1e6
    rug_mm = tramo.rugosidad_in * IN_A_MM
    res = ResultadoTramo(tramo.nombre, flujo_kbd, q_si, presion_recibo, dr, fr)
    if n == 0:
        return res

    J, K, L, M, P = [], [], [], [], []
    for p in tramo.puntos:
        d_in = p.diam_ext - 2 * p.espesor
        d_m = d_in * IN_A_M
        d_mm = d_in * IN_A_MM
        j = p.altura * densidad * PSI_POR_M_KGM3
        v = q_si / (0.25 * math.pi * d_m ** 2) if d_m > 0 else 0.0
        re_ = densidad * v * d_m / mu if mu > 0 else 0.0
        if tramo.formato_antiguo:
            # columna K de las plantillas antiguas (misma correlación de Churchill escrita con
            # constantes redondeadas; el %DR se aplica siempre)
            if q_si <= 0 or d_in <= 0:
                m = 0.0
            else:
                x = mu * d_in / (densidad * q_si)
                a = 3.1781e-20 * x ** 12
                b = 2.457 * math.log(0.17 * x ** 0.9 + 0.0106 * rug_mm / d_in)
                c = 9.746e45 * x ** 16
                m = ((800 / tramo.eficiencia) * (a + (b ** 16 + c) ** -1.5) ** (1 / 12) * densidad
                     * (q_si / (math.pi * d_m ** 2 / 4)) ** 2 / d_m * 1000 / 100000 / 2 * 14.5 * (1 - dr / 100))
        elif re_ <= 0:
            m = 0.0
        else:
            f = factor_friccion_churchill(re_, rug_mm, d_mm)
            m = ((100 / tramo.eficiencia) * f * densidad * v ** 2 / d_m * 1000 / 100000 / 2) * 14.5
            if re_ > RE_LAMINAR or tramo.dra_siempre:
                m *= (1 - dr / 100)
        maop = p.rating if (p.rating or 0) > 0 else FACTOR_DISENO * 2 * p.smys * p.espesor / p.diam_ext
        J.append(j); K.append(v); L.append(re_); M.append(m); P.append(maop)

    # Presión int+geod: se calcula de aguas abajo hacia aguas arriba
    N = [0.0] * n
    R = [False] * n                  # descuelgue: la línea no alcanza la cota del punto (columna separada)
    N[-1] = J[-1] + presion_recibo
    R[-1] = presion_recibo < 0
    for i in range(n - 2, -1, -1):
        e0, e1 = tramo.puntos[i].km, tramo.puntos[i + 1].km
        if e0 == e1:
            N[i] = N[i + 1]
            R[i] = R[i + 1]
        else:
            cand = N[i + 1] + (e1 - e0) * M[i]
            N[i] = cand if cand > J[i] else J[i]
            R[i] = cand < J[i] - 1e-9

    O = [N[i] - J[i] for i in range(n)]
    Q = [J[i] + P[i] for i in range(n)]
    S = [P[i] < O[i] * fr for i in range(n)]

    T = [0.0] * n
    for i in range(1, n):
        p = tramo.puntos[i]
        d_m = (p.diam_ext - 2 * p.espesor) * IN_A_M
        T[i] = (math.pi * d_m ** 2 / 4) * (p.km - tramo.puntos[i - 1].km) * 1e6 / 3.785 / 42 + T[i - 1]

    res.J, res.K, res.L, res.M, res.N, res.O = J, K, L, M, N, O
    res.P, res.Q, res.R, res.S, res.T = P, Q, R, S, T
    return res


def flujos_por_tramo(proy: Proyecto):
    """Resumen col E: flujo acumulado con inyecciones (bph*0.024 = kbd).
    Con ramales: el flujo que sigue por la línea principal después de cada derivación."""
    if getattr(proy, "ramales", None):
        return [r.flujo_kbd for r in calcular_red(proy)]
    flujos, prev = [], proy.flujo_kbd
    for est in proy.estaciones[:len(proy.tramos)]:
        f = prev
        if est.inyeccion == "INY 2":
            f = prev + proy.iny2_bph * BPH_A_KBD
        elif est.inyeccion == "INY 1":
            f = prev + proy.iny1_bph * BPH_A_KBD
        f += (getattr(est, "delta_bph", 0.0) or 0.0) * BPH_A_KBD
        flujos.append(f)
        prev = f
    return flujos


def calcular_proyecto(proy: Proyecto, flujo_kbd=None):
    """Calcula todos los tramos. Devuelve lista de ResultadoTramo (mismo orden que tramos).
    Si el sistema tiene ramales (derivaciones), devuelve una lista Resultados con el atributo .red."""
    if getattr(proy, "ramales", None):
        return calcular_red(proy, flujo_kbd)
    if flujo_kbd is not None:
        original = proy.flujo_kbd
        proy.flujo_kbd = flujo_kbd
        try:
            flujos = flujos_por_tramo(proy)
        finally:
            proy.flujo_kbd = original
    else:
        flujos = flujos_por_tramo(proy)

    n = len(proy.tramos)
    resultados = [None] * n
    # de aguas abajo hacia aguas arriba porque una estación "paso directo"
    # recibe a la presión de descarga del tramo siguiente
    for i in range(n - 1, -1, -1):
        est_ini = proy.estaciones[i]
        est_fin = proy.estaciones[i + 1]
        if est_fin.paso_directo and i + 1 < n and resultados[i + 1] is not None:
            p_rec = resultados[i + 1].presion_descarga
        else:
            p_rec = est_fin.presion_recibo
        resultados[i] = calcular_tramo(proy.tramos[i], flujos[i], proy.densidad, proy.viscosidad,
                                       est_ini.dr, est_ini.fr, p_rec)
    return resultados


def km_estaciones(proy: Proyecto):
    """Resumen col B: ubicación de cada estación."""
    kms = [0.0]
    for t in proy.tramos:
        kms.append(max((p.km for p in t.puntos), default=kms[-1]))
    if proy.tramos and proy.tramos[0].puntos:
        kms[0] = proy.tramos[0].puntos[0].km if proy.tramos[0].puntos[0].km else 0.0
    return kms


def interpolar_lista(proy: Proyecto, resultados, lista_km):
    """Tabla 'Interpolación de puntos requeridos' de la hoja Resumen para muchos km a la vez
    (búsqueda binaria: sirve para miles de puntos)."""
    from bisect import bisect_right
    kms = km_estaciones(proy)
    if not resultados or not proy.tramos:
        return [None for _ in lista_km]
    perfiles = [[p.km for p in t.puntos] for t in proy.tramos]
    alturas = [[p.altura for p in t.puntos] for t in proy.tramos]
    inicios = kms[:len(proy.tramos)]
    salida = []
    for km in lista_km:
        if km < kms[0] or km > kms[-1]:
            salida.append(None)
            continue
        idx = max(0, bisect_right(inicios, km) - 1)
        E = perfiles[idx]
        if len(E) < 2:
            salida.append(None)
            continue
        n = bisect_right(E, km) - 1                 # MATCH(km, E, 1)
        n = min(max(n, 0), len(E) - 2)
        while n > 0 and E[n + 1] == E[n]:
            n -= 1
        x0, x1 = E[n], E[n + 1]
        r = resultados[idx]

        def lin(arr):
            if x1 == x0:
                return arr[n]
            return arr[n] + (arr[n + 1] - arr[n]) * (km - x0) / (x1 - x0)

        salida.append({
            "km": km, "dist_tramo": km - kms[idx], "altura": lin(alturas[idx]), "p_geod": lin(r.J),
            "maop": lin(r.P), "p_int": lin(r.O), "maop_rem": lin(r.P) - lin(r.O), "p_int_geod": lin(r.N),
            "velocidad": lin(r.K), "nre": lin(r.L), "tramo": proy.tramos[idx].nombre,
            "inicio": proy.estaciones[idx].nombre, "fin": proy.estaciones[idx + 1].nombre,
            "fila": n + proy.tramos[idx].fila_inicio})
    return salida


def interpolar(proy: Proyecto, resultados, km):
    """Interpolación de un solo km (ver interpolar_lista)."""
    return interpolar_lista(proy, resultados, [km])[0]


def curva_sistema(proy: Proyecto, lim_inf, lim_sup, pasos=10, dr_forzado=None, cadena=""):
    """Macro CurvaSistema: presión de descarga de cada estación para flujos entre límites.

    dr_forzado: si no es None, se usa ese %DR en todos los tramos (p.ej. 0 = sin DRA).
    cadena: "" = línea principal (flujo = caudal que entra al sistema, con el reparto actual de los
        ramales); nombre de un ramal = curva propia del ramal (flujo = caudal que entra al ramal en su
        nodo; presión de descarga en el nodo hacia el ramal y en cada estación del ramal).
    """
    sistema = proy if not cadena else subred_ramal(proy, cadena)
    estaciones = list(sistema.estaciones) + [e for r in sistema.ramales for e in r.estaciones]
    originales = [e.dr for e in estaciones]
    try:
        if dr_forzado is not None:
            for e in estaciones:
                e.dr = dr_forzado
        filas = []
        for k in range(pasos + 1):
            q = lim_inf + (lim_sup - lim_inf) * k / pasos
            res = calcular_proyecto(sistema, flujo_kbd=q)
            filas.append((q, [r.presion_descarga for r in res]))
        return filas
    finally:
        for e, dr in zip(estaciones, originales):
            e.dr = dr




# --------------------------------------------------------------------------
# BUSCAR OBJETIVO (igual que Datos > Análisis de hipótesis > Buscar objetivo)
# --------------------------------------------------------------------------
#  Celda a definir (salida):  ("descarga", i) ("roturas", i|None) ("descuelgues", i) ("remanente", i)
#  Celda a cambiar (entrada): ("flujo", None) ("dr", i) ("recibo", j) ("iny1", None) ("iny2", None)
ES_CONTEO = ("roturas", "descuelgues", "roturas_r")
LIMITES_VAR = {"flujo": (0.0, 1e5), "bph": (0.0, 1e5 * KBD_A_BPH), "dr": (-1000.0, 1000.0), "recibo": (0.0, 1e5),
               "iny1": (-1e6, 1e6), "iny2": (-1e6, 1e6), "pct_ramal": (0.0, 100.0), "recibo_r": (0.0, 1e5)}
PASO_VAR = {"flujo": 1.0, "bph": 40.0, "dr": 2.0, "recibo": 5.0, "iny1": 50.0, "iny2": 50.0, "pct_ramal": 5.0,
            "recibo_r": 5.0}


def _ramal(proy, nombre):
    return next(r for r in proy.ramales if r.nombre == nombre)


def celdas_salida(proy: Proyecto):
    est, n = proy.estaciones, len(proy.tramos)
    out = [(f"Presión descarga – {est[i].nombre}", ("descarga", i)) for i in range(n)]
    out += [(f"Rotura de tubo – {est[i].nombre}", ("roturas", i)) for i in range(n)]
    out.append(("Rotura de tubo – TOTAL DEL SISTEMA", ("roturas", None)))
    out += [(f"Descuelgue – {est[i].nombre}", ("descuelgues", i)) for i in range(n)]
    out += [(f"Presión int remanente – {est[i].nombre}", ("remanente", i)) for i in range(n)]
    for r in getattr(proy, "ramales", []):
        out.append((f"Ramal {r.nombre}: flujo (KBDO)", ("flujo_ramal", r.nombre)))
        out.append((f"Ramal {r.nombre}: % del flujo que se deriva", ("pct_calc", r.nombre)))
        out.append((f"Ramal {r.nombre}: presión en el nodo", ("p_nodo", r.nombre)))
        out += [(f"Ramal {r.nombre}: presión descarga – {r.estaciones[i].nombre}", ("descarga_r", (r.nombre, i)))
                for i in range(len(r.tramos))]
        out.append((f"Ramal {r.nombre}: rotura de tubo (total)", ("roturas_r", r.nombre)))
    return out


def celdas_entrada(proy: Proyecto):
    est, n = proy.estaciones, len(proy.tramos)
    out = [("Flujo (KBDO)", ("flujo", None)), ("Flujo (bph) → ajusta el KBDO", ("bph", None))]
    out += [(f"%DR – {est[i].nombre}", ("dr", i)) for i in range(n)]
    out += [(f"Presión recibo – {est[j].nombre}", ("recibo", j)) for j in range(1, n + 1)
            if not est[j].paso_directo]
    out += [("INY 1 (bph)", ("iny1", None)), ("INY 2 (bph)", ("iny2", None))]
    for r in getattr(proy, "ramales", []):
        if r.modo != "Hidráulico":
            out.append((f"Ramal {r.nombre}: % del flujo", ("pct_ramal", r.nombre)))
        out += [(f"Ramal {r.nombre}: presión recibo – {r.estaciones[j].nombre}", ("recibo_r", (r.nombre, j)))
                for j in range(1, len(r.tramos) + 1) if not r.estaciones[j].paso_directo]
    return out


def leer_variable(proy: Proyecto, var):
    t, i = var
    return {"dr": lambda: proy.estaciones[i].dr, "flujo": lambda: proy.flujo_kbd,
            "bph": lambda: proy.flujo_kbd * KBD_A_BPH,
            "recibo": lambda: proy.estaciones[i].presion_recibo,
            "iny1": lambda: proy.iny1_bph, "iny2": lambda: proy.iny2_bph,
            "pct_ramal": lambda: _ramal(proy, i).pct,
            "recibo_r": lambda: _ramal(proy, i[0]).estaciones[i[1]].presion_recibo}[t]()


def fijar_variable(proy: Proyecto, var, valor):
    t, i = var
    if t == "dr":
        proy.estaciones[i].dr = valor
    elif t == "flujo":
        proy.flujo_kbd = valor
    elif t == "bph":
        proy.flujo_kbd = valor / KBD_A_BPH
    elif t == "recibo":
        proy.estaciones[i].presion_recibo = valor
    elif t == "iny1":
        proy.iny1_bph = valor
    elif t == "iny2":
        proy.iny2_bph = valor
    elif t == "pct_ramal":
        _ramal(proy, i).pct = valor
    elif t == "recibo_r":
        _ramal(proy, i[0]).estaciones[i[1]].presion_recibo = valor


def evaluar_celda(proy: Proyecto, celda, res=None):
    t, i = celda
    res = res if res is not None else calcular_proyecto(proy)
    if t == "descarga":
        return res[i].presion_descarga
    if t == "remanente":
        return res[i].presion_remanente or 0.0
    if t in ("flujo_ramal", "pct_calc", "p_nodo", "descarga_r", "roturas_r"):
        red = getattr(res, "red", None) or {}
        inf = red.get(i[0] if isinstance(i, tuple) else i)
        if inf is None:
            return 0.0
        if t == "flujo_ramal":
            return inf["flujo"]
        if t == "pct_calc":
            return inf["pct_calculado"]
        if t == "p_nodo":
            return inf["p_nodo"] or 0.0
        if t == "descarga_r":
            return inf["res"][i[1]].presion_descarga
        return sum(x.roturas for x in inf["res"])
    sel = res if i is None else [res[i]]
    if t == "roturas":
        return sum(r.roturas for r in sel)
    if t == "descuelgues":
        return sum(r.descuelgues for r in sel)
    raise ValueError(t)


@dataclass
class ResultadoBusqueda:
    encontrado: bool
    valor_variable: float
    valor_obtenido: float
    mensaje: str


def _expandir(x0, cambia, paso, lo, hi, max_pasos=60):
    """Avanza desde x0 hacia ambos lados (paso creciente) hasta que 'cambia(x)' sea True.
    Devuelve (x_donde_cambia, x_anterior_del_mismo_lado) o (None, None)."""
    h = paso
    prev = {1: x0, -1: x0}
    for _ in range(max_pasos):
        for s in (1, -1):
            x = min(max(x0 + s * h, lo), hi)
            if x == prev[s]:
                continue
            if cambia(x):
                return x, prev[s]
            prev[s] = x
        h *= 1.7
    return None, None


def buscar_objetivo(proy: Proyecto, celda, valor, variable, iteraciones=80) -> ResultadoBusqueda:
    """Busca el valor de 'variable' con el que 'celda' toma 'valor'. No modifica el proyecto.

    - Presiones: igual que Excel (la celda queda igual al valor).
    - Rotura de tubo / Descuelgue (conteos): se busca el valor LÍMITE más cercano al actual
      con el que el conteo es <= valor (p.ej. 0 roturas).
    """
    tipo = variable[0]
    lo, hi = LIMITES_VAR[tipo]
    x0 = leer_variable(proy, variable)
    paso = max(abs(x0) * 0.02, PASO_VAR[tipo])
    try:
        if celda[0] in ES_CONTEO:
            meta = int(round(valor))

            def ok(x):
                fijar_variable(proy, variable, x)
                return evaluar_celda(proy, celda) <= meta

            ok0 = ok(x0)
            x_c, x_prev = _expandir(x0, lambda x: ok(x) != ok0, paso, lo, hi)
            if x_c is None:
                if ok0:
                    return ResultadoBusqueda(True, x0, evaluar_celda(proy, celda),
                                             "La condición se cumple para cualquier valor de la variable.")
                return ResultadoBusqueda(False, x0, evaluar_celda(proy, celda),
                                         "Ningún valor de la variable cumple la condición: "
                                         "pruebe cambiando otra celda.")
            a, b = (x_prev, x_c) if ok0 else (x_c, x_prev)   # a cumple, b no cumple
            for _ in range(iteraciones):
                m = (a + b) / 2
                if ok(m):
                    a = m
                else:
                    b = m
            # se redondea a 4 decimales hacia el lado que cumple (valor "limpio" y seguro)
            x = math.floor(a * 1e4) / 1e4 if b > a else math.ceil(a * 1e4) / 1e4
            if not ok(x):
                x = a
            sentido = "máximo" if b > a else "mínimo"
            msg = f"Valor {sentido} de la variable con el que se cumple la condición (límite)."
        else:
            def f(x):
                fijar_variable(proy, variable, x)
                return evaluar_celda(proy, celda) - valor

            f0 = f(x0)
            if abs(f0) < 1e-6:
                x = x0
            else:
                x_c, x_prev = _expandir(x0, lambda x: (f(x) > 0) != (f0 > 0), paso, lo, hi)
                if x_c is None:
                    return ResultadoBusqueda(False, x0, f0 + valor,
                                             "No se encontró una solución: el valor objetivo no se "
                                             "alcanza cambiando esta celda.")
                a, b, fa = x_prev, x_c, f(x_prev)
                for _ in range(iteraciones):
                    m = (a + b) / 2
                    fm = f(m)
                    if abs(fm) < 1e-7:
                        a = b = m
                        break
                    if (fm > 0) == (fa > 0):
                        a, fa = m, fm
                    else:
                        b = m
                x = (a + b) / 2
            msg = "Se ha encontrado una solución."
        fijar_variable(proy, variable, x)
        return ResultadoBusqueda(True, x, evaluar_celda(proy, celda), msg)
    finally:
        fijar_variable(proy, variable, x0)


# --------------------------------------------------------------------------
# ACTUALIZACIÓN DE ESPESORES (nuevas inspecciones)
# --------------------------------------------------------------------------
TOL_KM = 1e-9


def _insertar_punto(tramo: Tramo, km):
    pts = tramo.puntos
    if any(abs(p.km - km) < TOL_KM for p in pts):
        return
    for i in range(len(pts) - 1):
        a, b = pts[i], pts[i + 1]
        if a.km < km < b.km:
            t = (km - a.km) / (b.km - a.km)
            pts.insert(i + 1, PuntoPerfil("", km, a.altura + (b.altura - a.altura) * t,
                                          a.espesor, a.diam_ext, a.smys, 0.0))
            return


def actualizar_espesor_rango(tramo: Tramo, km_ini, km_fin, espesor):
    """Asigna 'espesor' a la tubería entre km_ini y km_fin.

    Si los límites no existen en el perfil se insertan puntos (altura interpolada) y se
    duplican filas en los límites para representar el cambio de espesor en escalón,
    igual que en la plantilla Excel. Devuelve [(km, espesor_anterior, espesor_nuevo)].
    """
    if km_fin < km_ini:
        km_ini, km_fin = km_fin, km_ini
    if not tramo.puntos or km_fin <= km_ini:
        return []
    km_ini = max(km_ini, tramo.puntos[0].km)
    km_fin = min(km_fin, tramo.puntos[-1].km)
    if km_fin <= km_ini:
        return []
    _insertar_punto(tramo, km_ini)
    _insertar_punto(tramo, km_fin)
    pts = tramo.puntos

    def filas_en(km):
        return [i for i, p in enumerate(pts) if abs(p.km - km) < TOL_KM]

    ini = filas_en(km_ini)
    if len(ini) == 1 and ini[0] > 0:                       # fila que "llega" con el espesor viejo
        pts.insert(ini[0], replace(pts[ini[0]], punto="", rating=0.0))
    fin = filas_en(km_fin)
    if len(fin) == 1 and fin[0] < len(pts) - 1:            # fila que "sale" con el espesor viejo
        pts.insert(fin[0] + 1, replace(pts[fin[0]], punto="", rating=0.0))
    ini, fin = filas_en(km_ini), filas_en(km_fin)
    cambios = []
    for i, p in enumerate(pts):
        dentro = (km_ini < p.km < km_fin) or (i == ini[-1]) or (i == fin[0])
        if dentro and abs(p.espesor - espesor) > 1e-12:
            cambios.append((p.km, p.espesor, espesor))
            p.espesor = espesor
    return cambios


def leer_rangos_inspeccion(ruta):
    """Lee un archivo CSV o Excel con columnas: km_desde, km_hasta, espesor_in [, nota]."""
    filas = []
    if isinstance(ruta, tuple):                      # (nombre, bytes) de un archivo subido a la web
        nombre, contenido = ruta
    else:
        nombre = ruta
        with open(ruta, "rb") as fh:
            contenido = fh.read()
    if nombre.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=True).worksheets[0]
        crudas = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        texto = contenido.decode("utf-8-sig", errors="replace")
        dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
        crudas = list(csv.reader(io.StringIO(texto), dialecto))
    for r in crudas:
        if not r or len(r) < 3:
            continue
        try:
            a, b, e = (float(str(x).replace(",", ".")) for x in r[:3])
        except (TypeError, ValueError):
            continue          # encabezado o fila vacía
        nota = str(r[3]) if len(r) > 3 and r[3] is not None else ""
        filas.append((a, b, e, nota))
    return filas


# --------------------------------------------------------------------------
# PERFILES DE PRESIÓN EXTERNOS (MOP, transitorios / disparo de válvulas de SPS)
# y verificación contra la presión de diseño (Barlow / MAOP)
# --------------------------------------------------------------------------
TIPOS_PERFIL = ("MOP", "Transitorio", "Otro")


def _texto_norm(v):
    import unicodedata
    s = unicodedata.normalize("NFKD", str(v)).encode("ascii", "ignore").decode().lower()
    return " ".join(s.replace("\n", " ").split())


def leer_perfil_presion(origen, nombre_archivo=None):
    """Lee un perfil de presión interna (Excel o CSV) con columnas tipo 'Abscisa (Km)' y 'Presion Interna'.

    origen: ruta del archivo, o (nombre, bytes) de un archivo subido.
    Devuelve dict {nombre, tipo, archivo, km: [...], p: [...], fecha_carga}.
    """
    if isinstance(origen, tuple):
        nombre_archivo, contenido = origen
    else:
        nombre_archivo = nombre_archivo or os.path.basename(origen)
        with open(origen, "rb") as fh:
            contenido = fh.read()
    if nombre_archivo.lower().endswith((".xlsx", ".xlsm", ".xls")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
        ws = next(w for w in wb.worksheets if hasattr(w, "iter_rows"))
        filas = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        texto = contenido.decode("utf-8-sig", errors="replace")
        dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
        filas = [list(r) for r in csv.reader(io.StringIO(texto), dialecto)]

    # fila de encabezados: la que tiene una columna de km y una de presión interna
    c_km = c_p = fila_enc = None
    for i, fila in enumerate(filas[:60]):
        textos = {j: _texto_norm(v) for j, v in enumerate(fila) if isinstance(v, str) and v.strip()}
        km = next((j for j, t in textos.items() if t.startswith(("abscisa", "km", "dist", "ubicacion",
                                                                   "kilometr"))), None)
        pr = next((j for j, t in textos.items() if "presion interna" in t or t.startswith("presion int")), None)
        if pr is None:
            pr = next((j for j, t in textos.items() if t.startswith("presion") and "geod" not in t
                       and "diseno" not in t and "despacho" not in t and "recibo" not in t), None)
        if km is not None and pr is not None:
            c_km, c_p, fila_enc = km, pr, i
            break
    if fila_enc is None:
        raise ValueError("No encontré las columnas de abscisa (km) y presión interna en el archivo.")

    km, p = [], []
    for fila in filas[fila_enc + 1:]:
        if len(fila) <= max(c_km, c_p):
            continue
        a, b = fila[c_km], fila[c_p]
        try:
            a = float(str(a).replace(",", ".")) if not isinstance(a, (int, float)) else float(a)
            b = float(str(b).replace(",", ".")) if not isinstance(b, (int, float)) else float(b)
        except (TypeError, ValueError):
            continue
        km.append(a)
        p.append(b)
    if len(km) < 2:
        raise ValueError("El perfil tiene menos de 2 puntos con km y presión.")
    orden = sorted(range(len(km)), key=lambda k: km[k])
    km, p = [km[k] for k in orden], [p[k] for k in orden]

    # nombre y tipo: textos de las primeras filas (p.ej. 'SEBTOC-260226-MOP' / 'TRANSITORIO')
    textos = [str(v).strip() for fila in filas[:fila_enc] for v in fila if isinstance(v, str) and v.strip()]
    base = os.path.splitext(os.path.basename(nombre_archivo))[0]
    nombre = next((t for t in textos if "-" in t and " " not in t), base)
    todo = _texto_norm(" ".join(textos[:6]) + " " + base)
    tipo = "Transitorio" if ("transit" in todo or "disparo" in todo or "sps" in todo) else (
        "MOP" if "mop" in todo else "Otro")
    return {"nombre": nombre, "tipo": tipo, "archivo": os.path.basename(nombre_archivo), "km": km, "p": p,
            "fecha_carga": datetime.now().strftime("%Y-%m-%d %H:%M")}


def _barlow_en_km(proy, res, km_lista):
    """Presión de diseño (columna MAOP/Barlow del simulador) en cada km: la de la tubería en ese punto
    (si en ese km hay varias filas, la menor)."""
    from bisect import bisect_right, bisect_left
    kms = km_estaciones(proy)
    inicios = kms[:len(proy.tramos)]
    E = [[p.km for p in t.puntos] for t in proy.tramos]
    salida = []
    for x in km_lista:
        if not proy.tramos or x < kms[0] or x > kms[-1]:
            salida.append(None)
            continue
        idx = max(0, bisect_right(inicios, x) - 1)
        e, P = E[idx], res[idx].P
        a, b = bisect_left(e, x), bisect_right(e, x)
        if b > a:                                  # km exacto de una o varias filas
            salida.append(min(P[a:b]))
        else:
            salida.append(P[max(0, b - 1)])
    return salida


TOLERANCIA_DEFECTO = {"MOP": 0.0, "Transitorio": 10.0, "Otro": 0.0}


def tolerancia_perfil(perfil):
    """% que se permite superar la presión de diseño (transitorios: 10 % por defecto)."""
    t = perfil.get("tolerancia_pct")
    return float(t) if t is not None else TOLERANCIA_DEFECTO.get(perfil.get("tipo"), 0.0)


def verificar_perfil(proy, res, perfil, tolerancia_pct=None):
    """Compara un perfil de presión interna con la presión de diseño (Barlow / MAOP) del sistema.

    Evalúa en los km del perfil y en todos los puntos del simulador dentro del rango del perfil
    (interpolando el perfil). Para cada punto calcula el exceso en psi y en % de la presión de diseño:
        exceso % = (P perfil − P diseño) / P diseño × 100
    tolerancia_pct: % permitido por encima de la presión de diseño (p.ej. 10 % en transitorios).
    """
    from bisect import bisect_right
    tol = tolerancia_perfil(perfil) if tolerancia_pct is None else float(tolerancia_pct)
    kp, pp = perfil["km"], perfil["p"]

    def p_perfil(x):
        if x < kp[0] or x > kp[-1]:
            return None
        n = min(max(bisect_right(kp, x) - 1, 0), len(kp) - 2)
        x0, x1 = kp[n], kp[n + 1]
        return pp[n] if x1 == x0 else pp[n] + (pp[n + 1] - pp[n]) * (x - x0) / (x1 - x0)

    crudos = []
    for x, pr, bw in zip(kp, pp, _barlow_en_km(proy, res, kp)):
        if bw is not None:
            crudos.append((x, pr, bw, "perfil"))
    for t, r in zip(proy.tramos, res):
        for pt, bw in zip(t.puntos, r.P):
            pr = p_perfil(pt.km)
            if pr is not None:
                crudos.append((pt.km, pr, bw, pt.punto or "simulador"))
    crudos.sort(key=lambda z: z[0])
    # cada punto: (km, P perfil, P diseño, punto, exceso psi, exceso %, estado)
    puntos = []
    for x, pr, bw, nom in crudos:
        ex = pr - bw
        pct = 100.0 * ex / bw if bw else 0.0
        if ex <= 0:
            estado = "Cumple"
        elif pct <= tol + 1e-9:
            estado = f"Supera (dentro del {tol:g} %)"
        else:
            estado = f"Supera (fuera del {tol:g} %)" if tol > 0 else "Supera"
        puntos.append((x, pr, bw, nom, ex, pct, estado))
    excesos = [z for z in puntos if z[4] > 0]
    fuera = [z for z in excesos if z[5] > tol + 1e-9]

    # rangos continuos de km donde se supera la presión de diseño
    rangos, actual = [], None
    for z in puntos:
        if z[4] > 0:
            if actual is None:
                actual = {"km_ini": z[0], "km_fin": z[0], "max_pct": z[5], "km_max": z[0], "max_psi": z[4],
                          "n": 1, "fuera": z[5] > tol + 1e-9}
            else:
                actual["km_fin"] = z[0]
                actual["n"] += 1
                actual["fuera"] = actual["fuera"] or z[5] > tol + 1e-9
                if z[5] > actual["max_pct"]:
                    actual.update(max_pct=z[5], km_max=z[0], max_psi=z[4])
        elif actual is not None:
            rangos.append(actual)
            actual = None
    if actual is not None:
        rangos.append(actual)

    peor = max(puntos, key=lambda z: z[5]) if puntos else None
    if not excesos:
        resultado = "Cumple"
    elif not fuera and tol > 0:
        resultado = f"Supera la presión de diseño, dentro del {tol:g} % permitido"
    else:
        resultado = "Supera la presión de diseño" + (f" más del {tol:g} % permitido" if tol > 0 else "")
    return {
        "nombre": perfil["nombre"], "tipo": perfil["tipo"], "tolerancia_pct": tol, "n_puntos": len(puntos),
        "cumple": not excesos, "aceptable": not fuera, "resultado": resultado,
        "n_excesos": len(excesos), "n_fuera": len(fuera), "excesos": excesos, "rangos": rangos,
        "max_exceso_pct": peor[5] if peor else None, "max_exceso_psi": peor[4] if peor else None,
        "km_max_exceso": peor[0] if peor else None, "p_en_max": peor[1] if peor else None,
        "diseno_en_max": peor[2] if peor else None,
        "margen_min": -peor[4] if peor else None, "km_margen_min": peor[0] if peor else None,
        "p_en_margen_min": peor[1] if peor else None, "barlow_en_margen_min": peor[2] if peor else None,
        "uso_max_pct": (100 + peor[5]) if peor else None, "km_uso_max": peor[0] if peor else None,
        "puntos": puntos,
    }


# --------------------------------------------------------------------------
# PRESIÓN MÍNIMA DE RECIBO PARA EVITAR DESCUELGUE (columna separada / slack flow)
# --------------------------------------------------------------------------
def psi_a_m(psi, densidad):
    """Convierte presión (psi) a cabeza de líquido (m) con la densidad del producto."""
    k = densidad * PSI_POR_M_KGM3
    return psi / k if k else 0.0


def presion_minima_recibo(proy: Proyecto, res, margen_psi=None):
    """Para cada tramo calcula la presión de recibo mínima (en la estación de llegada) para que la línea
    no se descuelgue: la presión interna convertida a cabeza alcanza justo para pasar el punto más alto
    (presión interna = margen en el punto crítico y ≥ margen en todo el tramo).

    Con el caudal y %DR actuales la pérdida por fricción no depende de la presión de recibo, así que:
        N(i) = J(fin) + P_recibo + F(i)      F(i) = pérdida por fricción acumulada desde i hasta el recibo
        sin descuelgue  ⇔  N(i) − J(i) ≥ margen  en todos los puntos
        P_recibo mín = max_i [ J(i) + margen − F(i) ] − J(fin)
    Devuelve una lista (una entrada por tramo, i.e. por estación de recibo 1..n) de dicts.
    """
    margen = float(getattr(proy, "margen_descuelgue_psi", 0.0) or 0.0) if margen_psi is None else float(margen_psi)
    salida = []
    for i, (t, r) in enumerate(zip(proy.tramos, res)):
        n = len(t.puntos)
        est = proy.estaciones[i + 1] if i + 1 < len(proy.estaciones) else None
        d = {"tramo": t.nombre, "estacion": est.nombre if est else "", "bypass": bool(est and est.paso_directo),
             "margen_psi": margen, "p_min": None}
        if n == 0 or not r.J:
            salida.append(d)
            continue
        J, M = r.J, r.M
        F = [0.0] * n
        for k in range(n - 2, -1, -1):
            e0, e1 = t.puntos[k].km, t.puntos[k + 1].km
            F[k] = F[k + 1] if e0 == e1 else F[k + 1] + (e1 - e0) * M[k]
        req = [J[k] + margen - F[k] for k in range(n)]
        kc = max(range(n), key=lambda k: (req[k], -k))        # punto crítico (el primero si hay empate)
        p_min = req[kc] - J[-1]
        p_act = r.presion_recibo
        # puntos que hoy quedan en descuelgue (la presión interna calculada sería menor que el margen)
        N, desc = J[-1] + p_act, []
        for k in range(n - 2, -1, -1):
            e0, e1 = t.puntos[k].km, t.puntos[k + 1].km
            cand = N if e0 == e1 else N + (e1 - e0) * M[k]
            if cand - J[k] < margen - 1e-9:
                desc.append(k)
            N = max(cand, J[k])
        pc = t.puntos[kc]
        d.update({
            "p_min": p_min, "p_min_m": psi_a_m(p_min, proy.densidad),
            "p_actual": p_act, "p_actual_m": psi_a_m(p_act, proy.densidad), "holgura": p_act - p_min,
            "holgura_m": psi_a_m(p_act - p_min, proy.densidad), "descuelga": p_act < p_min - 1e-6,
            "km_critico": pc.km, "punto_critico": pc.punto, "altura_critica": pc.altura,
            "altura_recibo": t.puntos[-1].altura,
            "descarga_con_minimo": p_min + J[-1] + F[0] - J[0],
            "n_descuelgue": len(desc),
            "km_descuelgue": (min(t.puntos[k].km for k in desc), max(t.puntos[k].km for k in desc)) if desc else None,
            # línea int+geod con el recibo mínimo (toca el terreno en el punto crítico)
            "linea_minima": [J[-1] + p_min + F[k] for k in range(n)],
        })
        salida.append(d)
    return salida


# --------------------------------------------------------------------------
# ESCENARIOS: VÁLVULA ILÍCITA / DERRAME (fuga por un orificio en un km)
# --------------------------------------------------------------------------
PSI_A_PA = 6894.757
BBL_A_M3 = 0.158987294928
MODOS_FUGA = ("Caudal constante", "Presión de descarga constante")


def caudal_orificio_bph(diametro_in, dp_psi, densidad, cd=0.61):
    """Ecuación de orificio: q = Cd·A·√(2·ΔP/ρ). Devuelve bph."""
    if dp_psi <= 0 or diametro_in <= 0 or densidad <= 0:
        return 0.0
    area = math.pi / 4 * (diametro_in * IN_A_M) ** 2
    q_m3s = cd * area * math.sqrt(2 * dp_psi * PSI_A_PA / densidad)
    return q_m3s * 3600 / BBL_A_M3


def presion_orificio_psi(q_bph, diametro_in, densidad, cd=0.61):
    """Inversa: ΔP (psi) necesaria para que por el orificio pasen q_bph."""
    area = math.pi / 4 * (diametro_in * IN_A_M) ** 2
    if area <= 0 or cd <= 0:
        return 0.0
    q_m3s = q_bph * BBL_A_M3 / 3600
    return densidad / 2 * (q_m3s / (cd * area)) ** 2 / PSI_A_PA


def _raiz(f, a, b, fa=None, fb=None, tol=1e-6, maxit=60):
    """Regula falsi (Illinois) para f creciente o decreciente con cambio de signo en [a, b]."""
    fa = f(a) if fa is None else fa
    fb = f(b) if fb is None else fb
    if fa == 0:
        return a
    if fb == 0:
        return b
    lado = 0
    c = a
    for _ in range(maxit):
        c = (a * fb - b * fa) / (fb - fa) if fb != fa else (a + b) / 2
        fc = f(c)
        if abs(fc) < tol or abs(b - a) < tol * max(1.0, abs(c)):
            return c
        if (fc > 0) == (fb > 0):
            b, fb = c, fc
            if lado == -1:
                fa /= 2
            lado = -1
        else:
            a, fa = c, fc
            if lado == 1:
                fb /= 2
            lado = 1
    return c


def ubicar_tramo(cadena, km):
    """Índice del tramo de la cadena que contiene el km (None si está fuera)."""
    for i, t in enumerate(cadena.tramos):
        if t.puntos and t.puntos[0].km <= km <= t.puntos[-1].km:
            return i
    return None


def diametro_interno_en_km(cadena, km):
    i = ubicar_tramo(cadena, km)
    if i is None:
        return None
    pts = cadena.tramos[i].puntos
    j = max(k for k in range(len(pts)) if pts[k].km <= km)
    return pts[j].diam_ext - 2 * pts[j].espesor


def _partir_tramo(t, km):
    """Divide un tramo en el km: (aguas arriba, aguas abajo), con un punto interpolado en el km."""
    pts = t.puntos
    j = max(k for k in range(len(pts)) if pts[k].km <= km)
    a = pts[j]
    if a.km == km:
        nuevo = replace(a, punto="Fuga", rating=0.0)
    else:
        b = pts[min(j + 1, len(pts) - 1)]
        f = (km - a.km) / (b.km - a.km) if b.km != a.km else 0.0
        nuevo = replace(a, punto="Fuga", km=km, altura=a.altura + (b.altura - a.altura) * f, rating=0.0)
    arriba = [p for p in pts if p.km < km] + [nuevo]
    abajo = [nuevo] + [p for p in pts if p.km > km]
    if len(arriba) < 2:
        arriba = [pts[0], nuevo]
    if len(abajo) < 2:
        abajo = [nuevo, replace(nuevo, punto=pts[-1].punto)]
    return (replace(t, nombre=f"{t.nombre}a", puntos=arriba), replace(t, nombre=f"{t.nombre}b", puntos=abajo))


def simular_fuga(proy: Proyecto, km, diametro_in, cd=0.61, contrapresion_psi=0.0, modo="Caudal constante",
                 cadena="", tiempo_min=0.0):
    """Simula una fuga (válvula ilícita o derrame) por un orificio de 'diametro_in' en el km indicado.

    La fuga se representa como una estación virtual en el km (el tramo se parte en dos): aguas arriba
    pasa el caudal que llega y aguas abajo ese caudal menos la fuga. El caudal de fuga sale de la
    ecuación de orificio con la presión interna calculada en el punto (que a su vez baja con la fuga):
    se resuelve por iteración.
    modo "Caudal constante": el caudal que entra al sistema no cambia.
    modo "Presión de descarga constante": la estación de bombeo aguas arriba mantiene su presión de
        descarga y el caudal que entra aumenta (bombas de curva plana)."""
    import copy
    base_res = calcular_proyecto(proy)
    pc = copy.deepcopy(proy)
    pc.perfiles_externos = []
    cad = cadena_de(pc, cadena)
    if cad is None:
        raise ValueError("No se encontró la línea / ramal indicado.")
    idx = ubicar_tramo(cad, km)
    if idx is None:
        raise ValueError(f"El km {km} no está dentro de los tramos de esa línea.")
    t_arriba, t_abajo = _partir_tramo(cad.tramos[idx], km)
    est_ini = cad.estaciones[idx]
    virtual = Estacion(nombre=f"FUGA km {km:.3f}", paso_directo=True, dr=est_ini.dr, fr=est_ini.fr)
    cad.tramos[idx:idx + 1] = [t_arriba, t_abajo]
    cad.estaciones.insert(idx + 1, virtual)
    for r in pc.ramales:
        if (r.padre or "") == (cadena or "") and r.nodo > idx:
            r.nodo += 1

    def res_cadena(res):
        if not cadena:
            return res
        inf = (getattr(res, "red", None) or {}).get(cadena)
        return inf["res"] if inf else []

    densidad = pc.densidad
    estado = {"total": False}

    def evaluar(q_bph):
        virtual.delta_bph = -q_bph
        virtual.paso_directo = True
        res = calcular_proyecto(pc)
        rc = res_cadena(res)
        return res, rc

    def resolver_fuga():
        """Caudal de fuga con el flujo de entrada actual de pc."""
        res, rc = evaluar(0.0)
        q_llega = rc[idx].flujo_kbd / BPH_A_KBD
        px0 = rc[idx + 1].presion_descarga

        def f(q):
            _, rcq = evaluar(q)
            return q - caudal_orificio_bph(diametro_in, rcq[idx + 1].presion_descarga - contrapresion_psi, densidad, cd)
        f0 = -caudal_orificio_bph(diametro_in, px0 - contrapresion_psi, densidad, cd)
        if f0 >= 0:                                   # presión menor que la contrapresión: no sale nada
            estado["total"] = False
            return 0.0, q_llega
        f1 = f(q_llega)
        if f1 <= 0:
            # el orificio podría sacar más de lo que llega: sale todo y aguas abajo queda sin flujo
            estado["total"] = True
            return q_llega, q_llega
        estado["total"] = False
        return _raiz(f, 0.0, q_llega, f0, f1, tol=1e-6 * max(1.0, q_llega)), q_llega

    def aplicar(q, q_llega):
        res, rc = evaluar(q)
        if estado["total"]:
            virtual.paso_directo = False
            virtual.presion_recibo = contrapresion_psi + presion_orificio_psi(q_llega, diametro_in, densidad, cd)
            res = calcular_proyecto(pc)
            rc = res_cadena(res)
        return res, rc

    # estación de control (la que bombea aguas arriba de la fuga)
    s = idx
    while s > 0 and cad.estaciones[s].paso_directo:
        s -= 1
    rb = res_cadena(base_res)
    d_control_base = rb[s].presion_descarga if s < len(rb) else None
    flujo_base = proy.flujo_kbd
    nota_modo = ""
    if modo.startswith("Presión") and d_control_base is not None:
        def h(F):
            pc.flujo_kbd = F
            q, ql = resolver_fuga()
            _, rc = aplicar(q, ql)
            return rc[s].presion_descarga - d_control_base
        h0 = h(flujo_base)
        encontrado = abs(h0) < 1e-6
        if not encontrado:
            ant_F, ant_h = flujo_base, h0
            for mult in (1.05, 1.15, 1.3, 1.5, 2.0, 3.0):          # primero buscar hacia más caudal
                F = flujo_base * mult
                hF = h(F)
                if (hF > 0) != (ant_h > 0):
                    pc.flujo_kbd = _raiz(h, ant_F, F, ant_h, hF, tol=1e-4)
                    encontrado = True
                    break
                ant_F, ant_h = F, hF
        if not encontrado:
            ant_F, ant_h = flujo_base, h0
            for mult in (0.95, 0.85, 0.7, 0.5, 0.3):                # y si no, hacia menos caudal
                F = flujo_base * mult
                hF = h(F)
                if (hF > 0) != (ant_h > 0):
                    pc.flujo_kbd = _raiz(h, ant_F, F, ant_h, hF, tol=1e-4)
                    encontrado = True
                    break
                ant_F, ant_h = F, hF
        if not encontrado:
            pc.flujo_kbd = flujo_base
            nota_modo = ("No se encontró un caudal de entrada que mantenga la presión de descarga; "
                         "se usó caudal constante.")
    q, q_llega = resolver_fuga()
    res, rc = aplicar(q, q_llega)
    px = rc[idx + 1].presion_descarga if not estado["total"] else virtual.presion_recibo

    # presión interna en el km sin fuga (con la misma partición y fuga 0)
    pb = copy.deepcopy(pc)
    pb.flujo_kbd = flujo_base
    cb = cadena_de(pb, cadena)
    cb.estaciones[idx + 1].delta_bph = 0.0
    cb.estaciones[idx + 1].paso_directo = True
    res_b2 = calcular_proyecto(pb)
    rcb = res_cadena(res_b2)
    px_base = rcb[idx + 1].presion_descarga

    # comparación por estación de la cadena (sin la estación virtual)
    filas = []
    est_orig = cadena_de(proy, cadena).estaciones
    for i, e in enumerate(est_orig):
        j = i if i <= idx else i + 1                  # índice en la cadena con la fuga
        fila = {"Estación": e.nombre}
        if i < len(rb):
            fila["Descarga sin fuga (psig)"] = rb[i].presion_descarga
            fila["Flujo sin fuga (KBDO)"] = rb[i].flujo_kbd
        if j < len(rc):
            fila["Descarga con fuga (psig)"] = rc[j].presion_descarga
            fila["Flujo con fuga (KBDO)"] = rc[j].flujo_kbd
        if i > 0:
            fila["Recibo sin fuga (psig)"] = rb[i - 1].presion_recibo if i - 1 < len(rb) else None
            jr = (i - 1) if (i - 1) < idx else i
            if i - 1 == idx:
                jr = idx + 1
            fila["Recibo con fuga (psig)"] = rc[jr].presion_recibo if jr < len(rc) else None
            fila["Flujo que llega sin fuga (KBDO)"] = rb[i - 1].flujo_kbd if i - 1 < len(rb) else None
            fila["Flujo que llega con fuga (KBDO)"] = rc[jr].flujo_kbd if jr < len(rc) else None
        filas.append(fila)
    return {
        "km": km, "cadena": cadena, "tramo": cadena_de(proy, cadena).tramos[idx].nombre, "diametro_in": diametro_in,
        "cd": cd, "contrapresion_psi": contrapresion_psi, "modo": modo,
        "q_bph": q, "q_bpd": q * 24, "q_m3h": q * BBL_A_M3, "q_kbd": q * BPH_A_KBD,
        "volumen_bbl": q * tiempo_min / 60.0, "tiempo_min": tiempo_min,
        "px_base": px_base, "px_fuga": px, "flujo_llega_kbd": q_llega * BPH_A_KBD,
        "flujo_entrada_base": flujo_base, "flujo_entrada_fuga": pc.flujo_kbd,
        "estacion_control": est_orig[s].nombre, "d_control_base": d_control_base,
        "d_control_fuga": rc[s].presion_descarga if s < len(rc) else None,
        "toda_la_corriente": estado["total"], "pct_fuga": 100.0 * q / q_llega if q_llega else 0.0,
        "comparacion": filas, "proy_fuga": pc, "res_fuga": res, "res_base": base_res,
        "idx_tramo": idx, "nota": nota_modo,
    }


# --------------------------------------------------------------------------
# RED DE TRAMOS: ramales / derivaciones del flujo
# --------------------------------------------------------------------------
class Resultados(list):
    """Lista de ResultadoTramo de la línea principal + .red (resultados de los ramales y nodos)."""
    red = None


def _flujos_cadena(estaciones, n, q0, iny1, iny2, extraccion, ignorar_inicio=False):
    """Flujo de cada tramo de una cadena. extraccion[k] = fracción del flujo disponible en la estación k
    que se va por ramales. Devuelve (flujos por tramo, flujo disponible en cada estación 0..n)."""
    flujos, disp, prev = [], [], q0
    for k in range(n + 1):
        est = estaciones[k] if k < len(estaciones) else None
        f = prev
        if est is not None and k < n and not (ignorar_inicio and k == 0):
            if est.inyeccion == "INY 2":
                f += iny2 * BPH_A_KBD
            elif est.inyeccion == "INY 1":
                f += iny1 * BPH_A_KBD
            f += (getattr(est, "delta_bph", 0.0) or 0.0) * BPH_A_KBD
        disp.append(f)
        if k < n:
            f = f * (1 - extraccion.get(k, 0.0))
            flujos.append(f)
            prev = f
    return flujos, disp


def _presiones_cadena(estaciones, tramos, flujos, dens, visc, req):
    """Igual que calcular_proyecto (de aguas abajo hacia aguas arriba) pero en una estación en bypass con
    ramales, la presión del nodo es la mayor entre la que pide la línea y la que piden los ramales (req)."""
    n = len(tramos)
    res = [None] * n
    for i in range(n - 1, -1, -1):
        est_ini, est_fin = estaciones[i], estaciones[i + 1]
        if est_fin.paso_directo and i + 1 < n and res[i + 1] is not None:
            p_rec = res[i + 1].presion_descarga
            if i + 1 in req:
                p_rec = max(p_rec, req[i + 1])
        else:
            p_rec = est_fin.presion_recibo
        res[i] = calcular_tramo(tramos[i], flujos[i], dens, visc, est_ini.dr, est_ini.fr, p_rec)
    return res


def _hijos(proy, nombre_padre):
    return [r for r in proy.ramales if (r.padre or "") == nombre_padre]


def _evaluar(proy, cadena, nombre, q0, fracs, salida, es_ramal):
    """Calcula una cadena (línea principal o ramal) y, recursivamente, sus ramales."""
    n = len(cadena.tramos)
    hijos = [h for h in _hijos(proy, nombre) if 0 <= h.nodo <= n]
    ext = {}
    for h in hijos:
        ext[h.nodo] = ext.get(h.nodo, 0.0) + fracs[h.nombre]
    ext = {k: min(v, 1.0) for k, v in ext.items()}
    flujos, disp = _flujos_cadena(cadena.estaciones, n, q0, proy.iny1_bph, proy.iny2_bph, ext,
                                  ignorar_inicio=es_ramal)
    req = {}
    for h in hijos:
        qh = fracs[h.nombre] * disp[h.nodo]
        res_h = _evaluar(proy, h, h.nombre, qh, fracs, salida, True)
        d_h = res_h[0].presion_descarga if res_h else None
        salida[h.nombre] = {"ramal": h, "res": res_h, "flujo": qh, "fraccion": fracs[h.nombre],
                            "disponible_nodo": disp[h.nodo], "d_ramal": d_h, "padre": nombre, "nodo": h.nodo}
        if d_h is not None:
            req[h.nodo] = max(req.get(h.nodo, -1e99), d_h)
    res = _presiones_cadena(cadena.estaciones, cadena.tramos, flujos, proy.densidad, proy.viscosidad, req)
    for h in hijos:                                   # presión en el nodo y estrangulamiento
        info = salida[h.nombre]
        k = h.nodo
        d_padre = res[k].presion_descarga if k < n else None
        info["d_padre"] = d_padre
        candidatos = [x for x in (d_padre, req.get(k)) if x is not None]
        # presión en el nodo (cabezal común): la mayor que piden las ramas; si la estación está en bypass
        # esa presión se transmite aguas arriba, si bombea es la descarga que deben dar sus bombas
        p_nodo = max(candidatos) if candidatos else None
        info["p_nodo"] = p_nodo
        info["nodo_bombea"] = not (0 < k < n and cadena.estaciones[k].paso_directo)
        info["estrangula_ramal"] = (p_nodo - info["d_ramal"]) if (p_nodo is not None and info["d_ramal"] is not None) else None
        info["estrangula_padre"] = (p_nodo - d_padre) if (p_nodo is not None and d_padre is not None) else None
        info["flujo_padre_despues"] = (res[k].flujo_kbd if k < n else None)
    return res


def calcular_red(proy: Proyecto, flujo_kbd=None, iteraciones=40):
    """Calcula la línea principal y todos sus ramales.

    Ramales en modo Porcentaje: fracción fija. Ramales en modo Hidráulico: se busca la fracción con la
    que la presión que pide el ramal en el nodo es igual a la que pide la línea que continúa
    (mismo nodo, sin válvula de control); se resuelve por bisección (barridos sucesivos si hay varios)."""
    q0 = proy.flujo_kbd if flujo_kbd is None else flujo_kbd
    nombres = {r.nombre for r in proy.ramales}
    fracs = {r.nombre: min(max(r.pct / 100.0, 0.0), 1.0) for r in proy.ramales}
    hidraulicos = [r for r in proy.ramales if r.modo == "Hidráulico"]
    notas = {}

    def evaluar_todo():
        salida = {}
        res = _evaluar(proy, proy, "", q0, fracs, salida, False)
        return res, salida

    for _ in range(4 if len(hidraulicos) > 1 else 1):
        for h in hidraulicos:
            if h.padre and h.padre not in nombres:
                continue

            def g(fr):
                fracs[h.nombre] = fr
                res, sal = evaluar_todo()
                inf = sal.get(h.nombre)
                if not inf or inf["d_padre"] is None or inf["d_ramal"] is None:
                    return None
                return inf["d_padre"] - inf["d_ramal"]
            g0, g1 = g(0.0), g(1.0)
            if g0 is None:
                fracs[h.nombre] = min(max(h.pct / 100.0, 0.0), 1.0)
                notas[h.nombre] = ("El reparto hidráulico necesita que la línea continúe después del nodo; "
                                   "se usó el porcentaje.")
                continue
            if g0 <= 0:
                fracs[h.nombre] = 0.0
                notas[h.nombre] = "Con estas presiones de recibo el ramal no recibe flujo (la línea pide más presión)."
                continue
            if g1 >= 0:
                fracs[h.nombre] = 1.0
                notas[h.nombre] = "Con estas presiones de recibo todo el flujo se va por el ramal."
                continue
            fracs[h.nombre] = _raiz(g, 0.0, 1.0, g0, g1, tol=1e-9, maxit=iteraciones)
            notas.pop(h.nombre, None)
    res, salida = evaluar_todo()
    for nom, inf in salida.items():
        inf["nota"] = notas.get(nom, "")
        inf["pct_calculado"] = 100.0 * inf["fraccion"]
    out = Resultados(res)
    out.red = salida
    return out


def proyecto_ramal(proy: Proyecto, ramal, flujo=None):
    """Vista de un ramal como Proyecto (comparte estaciones y tramos) para graficar o exportar."""
    return Proyecto(sistema=f"{proy.sistema} – {ramal.nombre}", fecha=proy.fecha, responsable=proy.responsable,
                    flujo_kbd=flujo if flujo is not None else proy.flujo_kbd, producto=proy.producto,
                    densidad=proy.densidad, viscosidad=proy.viscosidad, iny1_bph=proy.iny1_bph,
                    iny2_bph=proy.iny2_bph, estaciones=ramal.estaciones, tramos=ramal.tramos,
                    productos=proy.productos, margen_descuelgue_psi=proy.margen_descuelgue_psi)


def cadena_de(proy, nombre):
    """Devuelve la cadena (Proyecto o Ramal) por nombre ('' = línea principal)."""
    if not nombre:
        return proy
    return next((r for r in proy.ramales if r.nombre == nombre), None)


def rutas_red(proy):
    """Rutas desde el origen hasta cada destino final: lista de (nombre, [ramales en orden])."""
    rutas = [("Línea principal: " + (proy.estaciones[0].nombre if proy.estaciones else "") + " → "
              + (proy.estaciones[-1].nombre if proy.estaciones else ""), [])]

    def camino(r):
        cad = [r]
        while cad[0].padre:
            p = cadena_de(proy, cad[0].padre)
            if p is None:
                break
            cad.insert(0, p)
        return cad
    for r in proy.ramales:
        cad = camino(r)
        ini = proy.estaciones[0].nombre if proy.estaciones else ""
        fin = r.estaciones[-1].nombre if r.estaciones else r.nombre
        rutas.append((f"{ini} → {fin} (por {' / '.join(x.nombre for x in cad)})", cad))
    return rutas


def proyecto_ruta(proy, res, cadena_ramales):
    """Une la línea principal hasta el nodo con los ramales del camino: (Proyecto, resultados) lineales
    para graficar el perfil hidráulico de la ruta. El km de los ramales se desplaza si no continúa el del
    nodo (p.ej. un ramal con km desde 0)."""
    if not cadena_ramales:
        return proy, list(res)
    red = getattr(res, "red", None) or {}
    est, tr, rs = [], [], []
    cad_actual, res_actual = proy, list(res)
    for r in cadena_ramales:
        k = r.nodo
        est += cad_actual.estaciones[:k]
        tr += cad_actual.tramos[:k]
        rs += res_actual[:k]
        inf = red.get(r.nombre)
        km_nodo = (max((p.km for p in cad_actual.tramos[k - 1].puntos), default=0.0) if k > 0
                   else (cad_actual.tramos[0].puntos[0].km if cad_actual.tramos and cad_actual.tramos[0].puntos else 0.0))
        if tr and k == 0:
            km_nodo = max(p.km for p in tr[-1].puntos)
        desplaz = 0.0
        if r.tramos and r.tramos[0].puntos:
            ini = r.tramos[0].puntos[0].km
            if abs(ini - km_nodo) > 0.5:
                desplaz = km_nodo - ini
        tramos_r = r.tramos
        if desplaz:
            tramos_r = [replace(t, puntos=[replace(p, km=p.km + desplaz) for p in t.puntos]) for t in r.tramos]
        est_r = list(r.estaciones)
        if est_r:
            est_r[0] = replace(est_r[0], nombre=cad_actual.estaciones[k].nombre if k < len(cad_actual.estaciones)
                               else est_r[0].nombre)
        cad_actual = Ramal(nombre=r.nombre, estaciones=est_r, tramos=tramos_r)
        res_actual = list(inf["res"]) if inf else []
    est += cad_actual.estaciones
    tr += cad_actual.tramos
    rs += res_actual
    ult = cadena_ramales[-1]
    p = Proyecto(sistema=f"{proy.sistema} – ruta {ult.nombre}", flujo_kbd=proy.flujo_kbd, producto=proy.producto,
                 densidad=proy.densidad, viscosidad=proy.viscosidad, estaciones=est, tramos=tr,
                 productos=proy.productos, margen_descuelgue_psi=proy.margen_descuelgue_psi,
                 perfiles_externos=[])
    return p, rs


def nuevo_ramal_desde_proyecto(proy: Proyecto, otro: Proyecto, nombre, padre="", nodo=0, modo="Porcentaje",
                               pct=50.0):
    """Crea un ramal con las estaciones y tramos de otro sistema (p.ej. importado de su plantilla Excel).
    La primera estación del otro sistema se toma como el nodo de derivación."""
    import copy
    est = copy.deepcopy(otro.estaciones)
    trs = copy.deepcopy(otro.tramos)
    cad = cadena_de(proy, padre)
    if cad is not None and est and nodo < len(cad.estaciones):
        est[0].nombre = cad.estaciones[nodo].nombre
    est[0].paso_directo = False
    return Ramal(nombre=nombre, padre=padre, nodo=nodo, modo=modo, pct=pct, estaciones=est, tramos=trs)


def nuevo_ramal_vacio(proy, nombre, padre, nodo, nombres_estaciones, modo="Porcentaje", pct=50.0):
    """Ramal con estaciones nuevas y tramos sin perfil (el perfil se pega o importa después)."""
    cad = cadena_de(proy, padre)
    nodo_nombre = cad.estaciones[nodo].nombre if cad is not None and nodo < len(cad.estaciones) else "Nodo"
    est = [Estacion(nombre=nodo_nombre)] + [Estacion(nombre=x) for x in nombres_estaciones]
    trs = [Tramo(nombre=f"{nombre}-{i + 1}") for i in range(len(nombres_estaciones))]
    return Ramal(nombre=nombre, padre=padre, nodo=nodo, modo=modo, pct=pct, estaciones=est, tramos=trs)


def posiciones_esquema(p):
    """Coordenadas del esquema de la red: línea principal en horizontal (y = 0); cada ramal en una fila
    nueva debajo, saliendo de su nodo. Devuelve {(cadena, índice estación): (x, y)}."""
    pos = {("", i): (float(i), 0.0) for i in range(len(p.estaciones))}
    fila = 0
    pendientes = list(p.ramales)
    while pendientes:
        # se ubica primero el ramal que sale más a la derecha: así cada fila nueva (más abajo) empieza a la
        # izquierda de las anteriores y su conector vertical no cruza estaciones de otros ramales
        listos = [r for r in pendientes if (r.padre or "", r.nodo) in pos]
        if not listos:
            break
        r = max(listos, key=lambda x: pos[(x.padre or "", x.nodo)][0])
        base = pos[(r.padre or "", r.nodo)]
        fila += 1
        y = -1.25 * fila
        for j in range(1, len(r.estaciones)):
            pos[(r.nombre, j)] = (base[0] + j, y)
        pos[(r.nombre, 0)] = base
        pendientes.remove(r)
    return pos


def subred_ramal(proy: Proyecto, nombre):
    """El ramal (y los ramales que salen de él) como un sistema propio alimentado en su nodo: sirve para la
    curva del sistema del ramal. Comparte estaciones y tramos con el sistema (salvo la del nodo)."""
    ram = cadena_de(proy, nombre)
    if ram is None or ram is proy:
        raise ValueError(f"No existe el ramal {nombre}.")
    nodo = replace(ram.estaciones[0], inyeccion="Sin inyección", delta_bph=0.0, paso_directo=False)
    desc, frontera = [], [ram.nombre]
    while frontera:
        hijos = [r for r in proy.ramales if r.padre in frontera]
        desc += hijos
        frontera = [h.nombre for h in hijos]
    subs = [replace(h, padre="" if h.padre == ram.nombre else h.padre) for h in desc]
    return Proyecto(sistema=f"{proy.sistema} – {ram.nombre}", flujo_kbd=proy.flujo_kbd, producto=proy.producto,
                    densidad=proy.densidad, viscosidad=proy.viscosidad, iny1_bph=proy.iny1_bph,
                    iny2_bph=proy.iny2_bph, estaciones=[nodo] + list(ram.estaciones[1:]), tramos=ram.tramos,
                    productos=proy.productos, ramales=subs)


def nombres_curva(proy: Proyecto, cadena=""):
    """Nombres de las estaciones (series) de la curva del sistema de una línea o ramal."""
    cad = proy if not cadena else cadena_de(proy, cadena)
    n = len(cad.tramos)
    nombres = [e.nombre for e in cad.estaciones[:n]]
    if cadena and nombres:
        nombres[0] = f"{nombres[0]} (hacia {cadena})"
    return nombres


def flujo_operacion(proy: Proyecto, res, cadena=""):
    """Caudal de operación actual de la línea (entrada al sistema) o del ramal (lo que entra en su nodo)."""
    if not cadena:
        return proy.flujo_kbd
    inf = (getattr(res, "red", None) or {}).get(cadena)
    return inf["flujo"] if inf else None

# --------------------------------------------------------------------------
# Importación desde la plantilla Excel
# --------------------------------------------------------------------------
def _num(v, defecto=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return defecto


def _nombre_hyperlink(v):
    if v is None:
        return ""
    s = str(v)
    if s.startswith("="):
        m = re.findall(r'"([^"]*)"', s)
        return m[-1] if m else s
    return s


def _abrir_libro(ruta, **kw):
    import openpyxl
    if isinstance(ruta, (bytes, bytearray)):
        return openpyxl.load_workbook(io.BytesIO(ruta), **kw)
    return openpyxl.load_workbook(ruta, **kw)


def _bloque(ws, max_fila, max_col, min_fila=1):
    """Lee un bloque de celdas de una hoja (modo rápido read_only) -> {(fila, col): valor}."""
    d = {}
    for r, fila in enumerate(ws.iter_rows(min_row=min_fila, max_row=max_fila, max_col=max_col, values_only=True),
                             min_fila):
        for c, v in enumerate(fila, 1):
            if v is not None:
                d[(r, c)] = v
    return d


def _nombre_hoja(valor, hojas):
    if valor in (None, ""):
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    s = str(valor).strip()
    if s in hojas:
        return s
    for h in hojas:                      # tolera espacios / mayúsculas distintas
        if h.strip().lower() == s.lower():
            return h
    return None


def _norm(v):
    """Texto en minúsculas, sin tildes ni espacios repetidos (para buscar títulos)."""
    import unicodedata
    s = unicodedata.normalize("NFKD", str(v)).encode("ascii", "ignore").decode().lower()
    return " ".join(s.replace("\n", " ").split())


def _es_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _buscar(celdas, condicion, max_fila=10 ** 6):
    """Primera celda (fila, col) cuyo texto cumple la condición (recorre por filas)."""
    for (r, c) in sorted(celdas):
        v = celdas[(r, c)]
        if r <= max_fila and isinstance(v, str) and condicion(_norm(v)):
            return r, c
    return None


def _num_derecha(celdas, pos, max_dist=4):
    """Primer número a la derecha de una etiqueta."""
    if not pos:
        return None
    r, c = pos
    for k in range(1, max_dist + 1):
        v = celdas.get((r, c + k))
        if _es_num(v):
            return float(v)
        if isinstance(v, str) and v.strip():
            return None                  # otra etiqueta: la celda de esta etiqueta está vacía
    return None


def importar_plantilla_excel(ruta, avance=None) -> Proyecto:
    """Lee una plantilla del simulador (.xlsm/.xlsx), versión nueva o antigua.

    Los datos se ubican por sus TÍTULOS ("Flujo", "%DR", "Inicio Segmento", "dist.del reg."...),
    no por posición fija, así funcionan las distintas versiones de la plantilla.
    ruta: ruta del archivo o su contenido en bytes. avance: función opcional avance(texto).
    Los avisos de la importación quedan en proy.avisos_importacion (lista de textos).
    """
    def aviso(t):
        if avance:
            avance(t)

    aviso("Abriendo el libro…")
    wb_v = _abrir_libro(ruta, data_only=True, read_only=True)     # valores calculados
    wb_f = _abrir_libro(ruta, data_only=False, read_only=True)    # fórmulas (nombres / bypass)
    hojas = [h for h in wb_v.sheetnames if hasattr(wb_v[h], "iter_rows")]   # sin hojas de gráfico
    if any(_norm(h).startswith("datos de entrada") for h in hojas) and \
            any(_norm(h).startswith("perfil hidraulico") for h in hojas):
        return _importar_formato_barlow(wb_v, wb_f, hojas, aviso)
    hoja_res = next((h for h in hojas if _norm(h) == "resumen"), None)
    if hoja_res is None:
        raise ValueError("El libro no tiene la hoja 'Resumen' (ni 'Datos de entrada' + 'Perfil hidráulico').")
    proy = Proyecto()
    avisos = []

    # ---------------- productos (hoja REF)
    hoja_ref = next((h for h in hojas if _norm(h) == "ref"), None)
    if hoja_ref:
        prods = []
        for fila in wb_v[hoja_ref].iter_rows(min_row=1, max_col=8, values_only=True):
            fila = list(fila) + [None] * (8 - len(fila))
            if fila[0] in (None, "") or not _es_num(fila[2]):
                continue
            prods.append([str(fila[0])] + [fila[c] for c in (2, 3, 4, 5, 6, 7)])
        if prods:
            proy.productos = prods
    nombres_prod = {_norm(p[0]): p for p in proy.productos}

    # ---------------- hoja Resumen
    aviso("Leyendo la hoja Resumen…")
    rv = _bloque(wb_v[hoja_res], 60, 25)
    rf = _bloque(wb_f[hoja_res], 60, 25)
    enc = _buscar(rv, lambda t: t.startswith("inicio segmento"))
    if not enc:
        raise ValueError("No encontré la tabla de estaciones ('Inicio Segmento') en la hoja Resumen.")
    hr, c0 = enc
    col = {}
    for (r, c), v in rv.items():
        if r == hr and isinstance(v, str):
            t = _norm(v)
            if t.startswith("ubicaci"):
                col["km"] = c
            elif "recibo" in t:
                col["recibo"] = c
            elif "descarga" in t:
                col["descarga"] = c
            elif t.startswith("flujo"):
                col["flujo"] = c
            elif t == "tramo":
                col["tramo"] = c
            elif t.startswith("inyecci"):
                col["iny"] = c

    sis = _buscar(rv, lambda t: t.startswith("sistema"), hr)
    if sis:
        proy.sistema = str(next((rv[(sis[0], c)] for c in range(sis[1] + 1, sis[1] + 6)
                                 if isinstance(rv.get((sis[0], c)), str)), "")).strip()
    fec = _buscar(rv, lambda t: t.startswith("fecha"), hr)
    if fec:
        v = next((rv[(fec[0], c)] for c in range(fec[1] + 1, fec[1] + 4) if rv.get((fec[0], c)) is not None), "")
        proy.fecha = v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else str(v)
    res_ = _buscar(rv, lambda t: t.startswith("responsable"), hr)
    if res_:
        proy.responsable = str(next((rv[(res_[0], c)] for c in range(res_[1] + 1, res_[1] + 6)
                                     if isinstance(rv.get((res_[0], c)), str)), ""))

    proy.flujo_kbd = _num_derecha(rv, _buscar(rv, lambda t: t.startswith("flujo"), hr - 1)) or 0.0
    proy.iny1_bph = _num_derecha(rv, _buscar(rv, lambda t: t.replace(" ", "") == "iny1", hr - 1)) or 0.0
    proy.iny2_bph = _num_derecha(rv, _buscar(rv, lambda t: t.replace(" ", "") == "iny2", hr - 1)) or 0.0

    # producto: a la derecha (formato antiguo) o debajo (formato nuevo) de la etiqueta "Producto(s)"
    pp = _buscar(rv, lambda t: t in ("producto", "productos"), hr - 1)
    producto = None
    if pp:
        for pos in ((pp[0], pp[1] + 1), (pp[0], pp[1] + 2), (pp[0] + 1, pp[1])):
            v = rv.get(pos)
            if isinstance(v, str) and _norm(v) in nombres_prod:
                producto = v
                break
    proy.producto = producto or "GM"
    proy.densidad = _num_derecha(rv, _buscar(rv, lambda t: t.startswith("densidad"), hr - 1), 6) or 0.0
    proy.viscosidad = _num_derecha(rv, _buscar(rv, lambda t: t.startswith("viscosidad"), hr - 1), 6) or 0.0
    p_ref = nombres_prod.get(_norm(proy.producto))
    if p_ref:
        proy.densidad = proy.densidad or _num(p_ref[1])
        proy.viscosidad = proy.viscosidad or _num(p_ref[2])

    # filas de %DR y FR: valores a la derecha de la etiqueta, uno por tramo
    def fila_valores(pos):
        if not pos:
            return []
        r, c = pos
        return [rv.get((r, c + 1 + k)) for k in range(20)]
    drs = fila_valores(_buscar(rv, lambda t: t in ("%dr", "%dra", "% dr", "% dra"), hr - 1))
    frs = fila_valores(_buscar(rv, lambda t: t == "fr", hr - 1))
    fr_comun = next((float(v) for v in frs if _es_num(v)), 1 / 1.1)

    # ---------------- estaciones y tramos
    filas_est = []
    r = hr + 1
    while r <= 60:
        v = rf.get((r, c0)) if rf.get((r, c0)) not in (None, "") else rv.get((r, c0))
        if v in (None, "") or _norm(v).startswith("total"):
            break
        filas_est.append(r)
        r += 1
    usadas = set()
    antiguo = False
    for k, fila in enumerate(filas_est):
        c_rec = col.get("recibo")
        c_f = rf.get((fila, c_rec)) if c_rec else None
        c_desc = col.get("descarga")
        paso = bool(c_desc) and isinstance(c_f, str) and \
            c_f.replace("$", "").strip().upper() == "=%s%d" % (_letra(c_desc), fila)
        dr = drs[k] if k < len(drs) and _es_num(drs[k]) else 0.0
        fr = frs[k] if k < len(frs) and _es_num(frs[k]) and frs[k] > 0 else fr_comun
        iny = str(rv.get((fila, col["iny"])) or "Sin inyección") if "iny" in col else "Sin inyección"
        est = Estacion(
            nombre=_nombre_hyperlink(rf.get((fila, c0))) or str(rv.get((fila, c0), f"EST {k + 1}")),
            presion_recibo=_num(rv.get((fila, c_rec)), 0) if c_rec else 0.0,
            paso_directo=paso, dr=float(dr), fr=float(fr),
            inyeccion=iny if iny in INYECCIONES else "Sin inyección")
        proy.estaciones.append(est)
        if k == len(filas_est) - 1:
            break                                         # la última estación no inicia tramo
        valor_h = rv.get((fila, col["tramo"])) if "tramo" in col else None
        hoja = _nombre_hoja(valor_h, hojas) or (_nombre_hoja(rf.get((fila, col["tramo"])), hojas)
                                                  if "tramo" in col else None)
        if hoja is None:
            hoja = _nombre_hoja(str(k + 1), hojas)        # hojas numeradas 1, 2, 3…
        if hoja is None:
            avisos.append(f"La estación {est.nombre} (fila {fila}) no tiene hoja de tramo válida.")
            break
        if hoja in usadas:
            avisos.append(f"La hoja '{hoja}' aparece repetida en la columna Tramo (fila {fila}).")
        usadas.add(hoja)
        aviso(f"Leyendo el tramo '{hoja}'…")
        tr, es_antiguo = _leer_tramo(wb_v[hoja], hoja)
        antiguo = antiguo or es_antiguo
        if not tr.puntos:
            avisos.append(f"La hoja '{hoja}' no tiene puntos de perfil.")
        proy.tramos.append(tr)
    proy.estaciones = proy.estaciones[:len(proy.tramos) + 1]
    for i in range(1, len(proy.tramos)):
        a, b = proy.tramos[i - 1].puntos, proy.tramos[i].puntos
        if a and b and abs(a[-1].km - b[0].km) > 0.5:
            avisos.append(f"El tramo '{proy.tramos[i - 1].nombre}' termina en km {a[-1].km:.3f} y el siguiente "
                          f"empieza en km {b[0].km:.3f}.")

    # ---------------- km de la tabla de interpolación
    km = _km_interpolacion(wb_v[hoja_res])
    if km:
        proy.km_interpolar = km
    proy.avisos_importacion = avisos
    proy.formato_antiguo = antiguo
    return proy


def _texto_derecha(celdas, pos, max_dist=6):
    if not pos:
        return None
    r, c = pos
    for k in range(1, max_dist + 1):
        v = celdas.get((r, c + k))
        if v not in (None, "") and not _es_num(v):
            return v
        if _es_num(v):
            return None
    return None


def _km_interpolacion(ws):
    """km de la tabla 'Interpolación de puntos requeridos' (columna bajo 'dist.del reg. [km]')."""
    pos = None
    for r, fila in enumerate(ws.iter_rows(min_row=1, max_row=400, max_col=30, values_only=True), 1):
        for c, v in enumerate(fila, 1):
            if isinstance(v, str) and _norm(v).startswith("dist.del reg"):
                pos = (r, c)
                break
        if pos:
            break
    if not pos:
        return []
    km, vacios = [], 0
    for (v,) in ws.iter_rows(min_row=pos[0] + 1, max_row=pos[0] + 5000, min_col=pos[1], max_col=pos[1],
                             values_only=True):
        if _es_num(v):
            km.append(float(v))
            vacios = 0
        else:
            vacios += 1
            if vacios > 5:
                break
    return km


def _importar_formato_barlow(wb_v, wb_f, hojas, aviso):
    """Plantilla nueva: hojas 'Datos de entrada', 'Perfil hidráulico' (un solo perfil continuo con la
    'Presión de diseño Barlow') y 'Propiedades fluido'. Caudales en bph por estación."""
    proy = Proyecto()
    avisos = []
    h_dat = next(h for h in hojas if _norm(h).startswith("datos de entrada"))
    h_per = next(h for h in hojas if _norm(h).startswith("perfil hidraulico"))
    h_flu = next((h for h in hojas if _norm(h).startswith("propiedades fluido")), None)

    if h_flu:
        filas = [list(r) for r in wb_v[h_flu].iter_rows(max_row=60, max_col=12, values_only=True)]
        enc = next(((i, j) for i, f in enumerate(filas) for j, v in enumerate(f)
                    if isinstance(v, str) and _norm(v) == "producto"), None)
        if enc:
            i0, j0 = enc
            cab = [_norm(v) if isinstance(v, str) else "" for v in filas[i0]]
            j_den = next((j for j, t in enumerate(cab) if t.startswith("densidad")), j0 + 1)
            j_vis = next((j for j, t in enumerate(cab) if t.startswith("viscosidad")), None)
            prods = []
            for f in filas[i0 + 1:]:
                f = f + [None] * (12 - len(f))
                if isinstance(f[j0], str) and f[j0].strip() and _es_num(f[j_den]) and \
                        j_vis is not None and _es_num(f[j_vis]):
                    prods.append([f[j0].strip(), float(f[j_den]), float(f[j_vis]), None, None, None, None])
            if prods:
                proy.productos = prods

    aviso("Leyendo 'Datos de entrada'…")
    dv = _bloque(wb_v[h_dat], 60, 30)
    df_ = _bloque(wb_f[h_dat], 60, 30)
    enc = _buscar(dv, lambda t: t.startswith("estacion") and "punto" in t)
    if not enc:
        raise ValueError("No encontré la tabla de estaciones en 'Datos de entrada'.")
    hr, c0 = enc
    proy.sistema = str(_texto_derecha(dv, _buscar(dv, lambda t: t == "sistema", hr)) or "").strip()
    proy.producto = str(_texto_derecha(dv, _buscar(dv, lambda t: t.startswith("fluido"), hr)) or "GM").strip()
    proy.densidad = _num_derecha(dv, _buscar(dv, lambda t: t.startswith("densidad"), hr), 3) or 0.0
    proy.viscosidad = _num_derecha(dv, _buscar(dv, lambda t: t.startswith("viscosidad cinem"), hr), 3) or 0.0
    rug = _num_derecha(dv, _buscar(dv, lambda t: t.startswith("rugosidad") and "( in" in t, hr), 3) or 0.0018
    for clave, attr in (("fecha", "fecha"), ("responsable", "responsable")):
        pos = _buscar(dv, lambda t, k=clave: t.startswith(k), hr)
        if pos:
            v = dv.get((pos[0] + 1, pos[1])) or _texto_derecha(dv, pos)
            if isinstance(v, str) and ("elaboraci" in _norm(v) or _norm(v).startswith(("fecha", "responsable"))):
                v = None                                  # es otra etiqueta, no el valor
            if v is not None:
                setattr(proy, attr, v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else str(v))
    p_ref = next((p for p in proy.productos if _norm(p[0]) == _norm(proy.producto)), None)
    if p_ref:
        proy.densidad = proy.densidad or float(p_ref[1])
        proy.viscosidad = proy.viscosidad or float(p_ref[2])

    col = {}
    for (r, c), v in dv.items():
        if r == hr and isinstance(v, str):
            t = _norm(v)
            if t.startswith("ubicaci"):
                col["km"] = c
            elif t.startswith("caudal despacho"):
                col["desp"] = c
            elif t.startswith("presion descarga"):
                col["desc"] = c
            elif t.startswith("presion recibo"):
                col["rec"] = c
            elif t.replace(" ", "") in ("%dr", "%dra"):
                col["dr"] = c
    filas_est = []
    r = hr + 1
    while dv.get((r, c0)) not in (None, ""):
        filas_est.append(r)
        r += 1
    despachos = [dv.get((f, col.get("desp"))) for f in filas_est]
    proy.flujo_kbd = float(despachos[0]) / KBD_A_BPH if despachos and _es_num(despachos[0]) else 0.0
    for k, f in enumerate(filas_est):
        rec_f = df_.get((f, col.get("rec")))
        paso = bool(col.get("desc")) and isinstance(rec_f, str) and \
            rec_f.replace("$", "").strip().upper() == "=%s%d" % (_letra(col["desc"]), f)
        dr = dv.get((f, col.get("dr")))
        est = Estacion(nombre=str(dv[(f, c0)]).strip(),
                       presion_recibo=_num(dv.get((f, col.get("rec"))), 0) if k > 0 else 0.0,
                       paso_directo=paso and k > 0, dr=float(dr) if _es_num(dr) else 0.0)
        if k > 0 and _es_num(despachos[k]) and _es_num(despachos[k - 1]):
            est.delta_bph = float(despachos[k]) - float(despachos[k - 1])
        proy.estaciones.append(est)

    # ---------------- perfil continuo -> tramos (según la fila de caudal que usa cada punto)
    aviso("Leyendo 'Perfil hidráulico'…")
    cab = _bloque(wb_v[h_per], 8, 40)
    pos = _buscar(cab, lambda t: t.startswith("abscisa"))
    if not pos:
        raise ValueError("No encontré la columna 'Abscisa' en 'Perfil hidráulico'.")
    hp = pos[0]
    cp = {}
    for (r, c), v in cab.items():
        if r == hp and isinstance(v, str):
            t = _norm(v)
            for clave, ini in (("sitio", "sitio"), ("km", "abscisa"), ("altura", "altura"),
                               ("diam", "diametro externo"), ("esp", "espesor"), ("smys", "smys"),
                               ("vel", "veloc")):
                if t.startswith(ini) and clave not in cp:
                    cp[clave] = c
    # fila del perfil donde arranca cada tramo: la fórmula de "Presión descarga" de cada estación
    # apunta a esa fila (p.ej. ='Perfil hidráulico'!X318)
    inicio_fila = {}
    for k, f in enumerate(filas_est[:-1]):
        fo = df_.get((f, col.get("desc"))) if "desc" in col else None
        m_ = re.search(r"!\$?[A-Z]+\$?(\d+)", fo) if isinstance(fo, str) else None
        if m_:
            inicio_fila[k] = int(m_.group(1))
    patron_q = re.compile(r"\*\s*([0-9]*\.?[0-9]+(?:[Ee]-?\d+)?)\s*\)")
    factor_bph = 0.00004416
    if "vel" in cp:
        for (fo,) in wb_f[h_per].iter_rows(min_row=hp + 1, max_row=hp + 30, min_col=cp["vel"], max_col=cp["vel"],
                                           values_only=True):
            mq = patron_q.search(fo) if isinstance(fo, str) else None
            if mq:
                factor_bph = float(mq.group(1))
                break
    ultima = max(cp.values())
    filas_perfil = []
    for i, fila in enumerate(wb_v[h_per].iter_rows(min_row=hp + 1, max_col=ultima, values_only=True)):
        fila = list(fila) + [None] * (ultima - len(fila))

        def v(clave):
            return fila[cp[clave] - 1] if clave in cp else None
        km = v("km")
        if not _es_num(km):
            if km is None and filas_perfil:
                break
            continue
        filas_perfil.append((hp + 1 + i, PuntoPerfil(str(v("sitio") or ""), float(km), _num(v("altura")),
                                                     _num(v("esp")), _num(v("diam")), _num(v("smys")), 0.0)))
    n_tr = len(filas_est) - 1
    if len(inicio_fila) == n_tr and n_tr > 0:
        inicios = [inicio_fila[k] for k in range(n_tr)]
        inicios[0] = min(inicios[0], filas_perfil[0][0]) if filas_perfil else inicios[0]
    else:                                           # sin fórmulas: se corta por el km de cada estación
        kms_est = [_num(dv.get((r_, col.get("km")))) for r_ in filas_est]
        inicios = [filas_perfil[0][0]] if filas_perfil else []
        for k in range(1, n_tr):
            fila_k = next((fr for fr, pt in filas_perfil if pt.km >= kms_est[k] and fr > inicios[-1]), None)
            if fila_k is None:
                break
            iguales = [fr for fr, pt in filas_perfil if abs(pt.km - kms_est[k]) < 1e-9]
            inicios.append(iguales[-1] if len(iguales) > 1 else fila_k)
        avisos.append("Los tramos se separaron por el km de cada estación (no se encontraron las fórmulas "
                      "de presión de descarga).")
    for k, ini in enumerate(inicios):
        fin = inicios[k + 1] if k + 1 < len(inicios) else 10 ** 9
        pts = [pt for fr, pt in filas_perfil if ini <= fr < fin]
        if not pts:
            continue
        t = Tramo(nombre=str(k + 1), rugosidad_in=rug, eficiencia=100.0, dra_siempre=True,
                  factor_q=factor_bph * KBD_A_BPH, fila_inicio=ini)
        t.puntos = pts
        proy.tramos.append(t)
    if len(proy.tramos) != len(proy.estaciones) - 1:
        avisos.append(f"Se encontraron {len(proy.tramos)} tramos para {len(proy.estaciones)} estaciones.")
        proy.estaciones = proy.estaciones[:len(proy.tramos) + 1]
    km_int = _km_interpolacion(wb_v[h_dat])
    if km_int:
        proy.km_interpolar = km_int
    proy.avisos_importacion = avisos
    proy.formato_antiguo = False
    return proy


def _letra(c):
    from openpyxl.utils import get_column_letter
    return get_column_letter(c)


def _leer_tramo(ws, nombre):
    """Lee una hoja de tramo. Devuelve (Tramo, es_formato_antiguo)."""
    cab = _bloque(ws, 6, 30)
    pos = _buscar(cab, lambda t: t.startswith("dist.del reg") or t.startswith("dist. del reg"))
    if not pos:
        raise ValueError(f"La hoja '{nombre}' no tiene la columna 'dist.del reg. [km]'.")
    hr = pos[0]
    col = {}
    for (r, c), v in cab.items():
        if r == hr and isinstance(v, str):
            t = _norm(v)
            if t.startswith("punto"):
                col["punto"] = c
            elif t.startswith("dist"):
                col["km"] = c
            elif t.startswith("altura"):
                col["altura"] = c
            elif t.startswith("esp"):
                col["esp"] = c
            elif t.startswith("diam"):
                col["diam"] = c
            elif t.startswith("smys"):
                col["smys"] = c
            elif t.startswith("rating"):
                col["rating"] = c
            elif t == "nre":
                col["nre"] = c
    for k in ("altura", "esp", "diam", "smys"):
        if k not in col:
            raise ValueError(f"La hoja '{nombre}' no tiene la columna de {k}.")
    antiguo = "nre" not in col                 # las plantillas antiguas no tienen Reynolds

    # parámetros de la columna A/B: rugosidad (in) y eficiencia
    par = _bloque(ws, 40, 2)
    rug, eff = 0.0018, 100.0
    for (r, c), v in par.items():
        if c == 1 and isinstance(v, str):
            t = _norm(v)
            if t.startswith("rugosidad") and "( in" in t.replace("(in", "( in") and _es_num(par.get((r, 2))):
                rug = float(par[(r, 2)])
            elif t.startswith("eficiencia") and _es_num(par.get((r, 2))):
                eff = float(par[(r, 2)]) or 100.0
    t = Tramo(nombre=nombre, rugosidad_in=rug, eficiencia=eff, formato_antiguo=antiguo, fila_inicio=hr + 1)
    ultima = max(col.values())
    for row in ws.iter_rows(min_row=hr + 1, max_col=ultima, values_only=True):
        row = list(row) + [None] * (ultima - len(row))

        def v(clave, defecto=None):
            c = col.get(clave)
            return row[c - 1] if c else defecto
        km = v("km")
        if km is None or km == "":
            break
        if not _es_num(km) or not _es_num(v("altura")):
            continue
        t.puntos.append(PuntoPerfil(str(v("punto") or ""), float(km), _num(v("altura")), _num(v("esp")),
                                    _num(v("diam")), _num(v("smys")), _num(v("rating"), 0)))
    return t, antiguo


# --------------------------------------------------------------------------
# Perfil desde CSV
# --------------------------------------------------------------------------
COLUMNAS_CSV = ["punto", "km", "altura_m", "espesor_in", "diam_ext_in", "smys_psi", "rating_psi"]


def leer_perfil_csv(ruta, nombre) -> Tramo:
    with open(ruta, newline="", encoding="utf-8-sig") as fh:
        muestra = fh.read(4096)
        fh.seek(0)
        dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t")
        filas = list(csv.reader(fh, dialecto))
    t = Tramo(nombre=nombre)
    for fila in filas:
        if not fila:
            continue
        try:
            vals = [float(str(x).replace(",", ".")) if x.strip() else 0.0 for x in fila[1:7]]
        except ValueError:
            continue  # encabezado
        while len(vals) < 6:
            vals.append(0.0)
        t.puntos.append(PuntoPerfil(fila[0], *vals))
    return t


def escribir_perfil_csv(tramo: Tramo, ruta):
    with open(ruta, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(COLUMNAS_CSV)
        for p in tramo.puntos:
            w.writerow([p.punto, p.km, p.altura, p.espesor, p.diam_ext, p.smys, p.rating])


# ==========================================================================
# BASE DE DATOS DE SISTEMAS (SQLite)
# ==========================================================================
def carpeta_programa():
    try:
        return os.path.dirname(os.path.abspath(__file__))
    except NameError:          # ejecutado desde una consola interactiva
        return os.getcwd()


# Se puede cambiar con la variable de entorno SIMULADOR_BD (p.ej. una carpeta del servidor)
RUTA_BD = os.environ.get("SIMULADOR_BD", os.path.join(carpeta_programa(), "simuladores.db"))

CAMPOS_PUNTO = [("punto", "Punto/ubicación"), ("km", "km"), ("altura", "Altura (m)"),
                ("espesor", "Espesor (in)"), ("diam_ext", "Diámetro ext. (in)"),
                ("smys", "SMYS (psi)"), ("rating", "Rating válvula (psi)")]


def usuario_actual():
    try:
        return getpass.getuser()
    except Exception:  # noqa: BLE001
        return "usuario"


class BaseDatos:
    """sistemas: un registro por sistema con todos sus datos (JSON).
    historial: trazabilidad de cambios (espesores, rating, operación…).
    productos: tabla REF compartida por todos los sistemas."""

    def __init__(self, ruta=RUTA_BD):
        self.ruta = ruta
        self.con = sqlite3.connect(ruta, check_same_thread=False, timeout=30)
        self.con.execute("PRAGMA journal_mode=WAL")      # varios usuarios leyendo a la vez
        c = self.con
        c.execute("CREATE TABLE IF NOT EXISTS sistemas (nombre TEXT PRIMARY KEY, datos TEXT NOT NULL, "
                  "actualizado TEXT, usuario TEXT)")
        c.execute("CREATE TABLE IF NOT EXISTS historial (id INTEGER PRIMARY KEY AUTOINCREMENT, sistema TEXT, "
                  "fecha TEXT, usuario TEXT, tramo TEXT, km REAL, campo TEXT, anterior TEXT, nuevo TEXT, nota TEXT)")
        c.execute("CREATE TABLE IF NOT EXISTS productos (orden INTEGER, nombre TEXT PRIMARY KEY, densidad REAL, "
                  "visc1 REAL, t1 REAL, dens2 REAL, visc2 REAL, t2 REAL)")
        c.execute("CREATE TABLE IF NOT EXISTS config (clave TEXT PRIMARY KEY, valor TEXT)")
        c.commit()
        if not self.productos():
            self.guardar_productos(PRODUCTOS_DEFECTO)

    # ---------------- sistemas
    def sistemas(self):
        return [r[0] for r in self.con.execute("SELECT nombre FROM sistemas ORDER BY nombre COLLATE NOCASE")]

    def existe(self, nombre):
        return self.con.execute("SELECT 1 FROM sistemas WHERE nombre=?", (nombre,)).fetchone() is not None

    def cargar(self, nombre) -> Proyecto:
        fila = self.con.execute("SELECT datos FROM sistemas WHERE nombre=?", (nombre,)).fetchone()
        if not fila:
            raise KeyError(nombre)
        p = Proyecto.from_dict(json.loads(fila[0]))
        p.sistema = nombre
        p.productos = self.productos()
        return p

    def info(self, nombre):
        return self.con.execute("SELECT actualizado, usuario FROM sistemas WHERE nombre=?", (nombre,)).fetchone()

    def guardar(self, proy: Proyecto, nota="", cambios_extra=None, usuario=None):
        """Guarda el sistema y registra en el historial lo que cambió. Devuelve # de cambios."""
        ahora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        usuario = usuario or usuario_actual()
        if self.existe(proy.sistema):
            cambios = diferencias(self.cargar(proy.sistema), proy)
        else:
            cambios = [("", None, "Sistema", "", "Creado en la base de datos")]
        extra = [tuple(c) + ("",) * (6 - len(c)) for c in (cambios_extra or [])]
        ya = {(c[0], c[1], c[2]) for c in extra}
        cambios = [tuple(c) + (nota,) for c in cambios if (c[0], c[1], c[2]) not in ya]
        cambios += [c[:5] + (c[5] or nota,) for c in extra]
        datos = proy.to_dict()
        datos.pop("productos", None)
        self.con.execute("INSERT OR REPLACE INTO sistemas VALUES (?,?,?,?)",
                         (proy.sistema, json.dumps(datos, ensure_ascii=False), ahora, usuario))
        self.con.executemany(
            "INSERT INTO historial (sistema, fecha, usuario, tramo, km, campo, anterior, nuevo, nota) "
            "VALUES (?,?,?,?,?,?,?,?,?)",
            [(proy.sistema, ahora, usuario, t, km, campo, str(a), str(n), nt) for t, km, campo, a, n, nt in cambios])
        self.guardar_productos(proy.productos, commit=False)
        self.con.commit()
        return len(cambios)

    def eliminar(self, nombre):
        self.con.execute("DELETE FROM sistemas WHERE nombre=?", (nombre,))
        self.con.commit()

    def renombrar(self, viejo, nuevo):
        self.con.execute("UPDATE sistemas SET nombre=? WHERE nombre=?", (nuevo, viejo))
        self.con.execute("UPDATE historial SET sistema=? WHERE sistema=?", (nuevo, viejo))
        self.con.commit()

    def historial(self, nombre):
        return self.con.execute("SELECT fecha, usuario, tramo, km, campo, anterior, nuevo, nota FROM historial "
                                "WHERE sistema=? ORDER BY id DESC", (nombre,)).fetchall()

    # ---------------- productos
    def productos(self):
        return [list(r) for r in self.con.execute(
            "SELECT nombre, densidad, visc1, t1, dens2, visc2, t2 FROM productos ORDER BY orden")]

    def guardar_productos(self, productos, commit=True):
        self.con.execute("DELETE FROM productos")
        self.con.executemany("INSERT OR REPLACE INTO productos VALUES (?,?,?,?,?,?,?,?)",
                             [(k, *(list(p) + [None] * 7)[:7]) for k, p in enumerate(productos)])
        if commit:
            self.con.commit()

    # ---------------- configuración
    def config(self, clave, defecto=None):
        r = self.con.execute("SELECT valor FROM config WHERE clave=?", (clave,)).fetchone()
        return r[0] if r else defecto

    def fijar_config(self, clave, valor):
        self.con.execute("INSERT OR REPLACE INTO config VALUES (?,?)", (clave, str(valor)))
        self.con.commit()


def _fmt_hist(v):
    if isinstance(v, float):
        return f"{v:.10g}"
    return str(v)


def diferencias(viejo: Proyecto, nuevo: Proyecto):
    """Lista de cambios (tramo, km, campo, anterior, nuevo) entre dos versiones de un sistema."""
    cambios = []
    for campo, et in [("flujo_kbd", "Flujo (KBDO)"), ("producto", "Producto"), ("iny1_bph", "INY 1 (bph)"),
                      ("iny2_bph", "INY 2 (bph)"), ("densidad", "Densidad"), ("viscosidad", "Viscosidad"),
                      ("margen_descuelgue_psi", "Margen contra descuelgue (psi)")]:
        a, b = getattr(viejo, campo, 0.0), getattr(nuevo, campo, 0.0)
        if a != b:
            cambios.append(("", None, et, _fmt_hist(a), _fmt_hist(b)))
    for eo, en in zip(viejo.estaciones, nuevo.estaciones):
        for campo, et in [("nombre", "Nombre"), ("presion_recibo", "Presión recibo"), ("paso_directo", "Bypass"),
                          ("dr", "%DR"), ("fr", "FR"), ("inyeccion", "Inyección"),
                          ("delta_bph", "Iny(+)/Entrega(−) bph")]:
            a, b = getattr(eo, campo, 0.0), getattr(en, campo, 0.0)
            if campo == "presion_recibo" and en.paso_directo:
                continue
            if a != b:
                cambios.append((f"Estación {en.nombre}", None, et, _fmt_hist(a), _fmt_hist(b)))
    rv = {r.nombre: r for r in getattr(viejo, "ramales", [])}
    rn = {r.nombre: r for r in getattr(nuevo, "ramales", [])}
    for nom in rn.keys() - rv.keys():
        cambios.append((f"Ramal {nom}", None, "Ramal agregado", "", f"{len(rn[nom].tramos)} tramos"))
    for nom in rv.keys() - rn.keys():
        cambios.append((f"Ramal {nom}", None, "Ramal quitado", f"{len(rv[nom].tramos)} tramos", ""))
    for nom in rn.keys() & rv.keys():
        a_, b_ = rv[nom], rn[nom]
        for campo, et in (("modo", "Modo de reparto"), ("pct", "% del flujo"), ("nodo", "Nodo"), ("padre", "Padre")):
            if getattr(a_, campo) != getattr(b_, campo):
                cambios.append((f"Ramal {nom}", None, et, _fmt_hist(getattr(a_, campo)), _fmt_hist(getattr(b_, campo))))
        for eo, en in zip(a_.estaciones, b_.estaciones):
            for campo, et in [("presion_recibo", "Presión recibo"), ("paso_directo", "Bypass"), ("dr", "%DR"),
                              ("fr", "FR"), ("delta_bph", "Iny(+)/Entrega(−) bph")]:
                x, y = getattr(eo, campo, 0.0), getattr(en, campo, 0.0)
                if x != y:
                    cambios.append((f"Ramal {nom} – {en.nombre}", None, et, _fmt_hist(x), _fmt_hist(y)))
        for to, tn in zip(a_.tramos, b_.tramos):
            if len(to.puntos) != len(tn.puntos):
                cambios.append((f"Ramal {nom} – {tn.nombre}", None, "Puntos del perfil", len(to.puntos), len(tn.puntos)))
                continue
            for po, pn in zip(to.puntos, tn.puntos):
                for campo, et in CAMPOS_PUNTO:
                    x, y = getattr(po, campo), getattr(pn, campo)
                    if x != y:
                        cambios.append((f"Ramal {nom} – {tn.nombre}", pn.km, et, _fmt_hist(x), _fmt_hist(y)))
    if len(viejo.tramos) != len(nuevo.tramos):
        cambios.append(("", None, "Número de tramos", len(viejo.tramos), len(nuevo.tramos)))
    for to, tn in zip(viejo.tramos, nuevo.tramos):
        if to.rugosidad_in != tn.rugosidad_in:
            cambios.append((tn.nombre, None, "Rugosidad (in)", _fmt_hist(to.rugosidad_in), _fmt_hist(tn.rugosidad_in)))
        if to.eficiencia != tn.eficiencia:
            cambios.append((tn.nombre, None, "Eficiencia (%)", _fmt_hist(to.eficiencia), _fmt_hist(tn.eficiencia)))
        if len(to.puntos) != len(tn.puntos):
            cambios.append((tn.nombre, None, "Puntos del perfil", len(to.puntos), len(tn.puntos)))
            continue
        for po, pn in zip(to.puntos, tn.puntos):
            for campo, et in CAMPOS_PUNTO:
                a, b = getattr(po, campo), getattr(pn, campo)
                if a != b:
                    cambios.append((tn.nombre, pn.km, et, _fmt_hist(a), _fmt_hist(b)))
    return cambios
# ==========================================================================
# Exportación a Excel
# ==========================================================================
def exportar_resultados_excel(proy, res, ruta, curva=None, curva_nombres=None, curva_titulo=None):
    """ruta: nombre de archivo o io.BytesIO (descarga desde la web).
    curva_nombres / curva_titulo: series de la curva del sistema (p.ej. la curva de un ramal)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    neg = Font(bold=True, color="FFFFFF")
    azul = PatternFill("solid", fgColor="1F4E79")
    rojo = PatternFill("solid", fgColor="FFC7CE")
    amar = PatternFill("solid", fgColor="FFEB9C")

    def encabezado(ws, fila, textos):
        for j, t in enumerate(textos, 1):
            c = ws.cell(fila, j, t)
            c.font, c.fill = neg, azul
            c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(j)].width = max(12, min(30, len(t) + 2))

    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "SIMULADOR EN ESTADO ESTABLE PARA LA GENERACIÓN DE PERFILES HIDRÁULICOS PARA PRIORIZACIÓN ILI"
    ws["A1"].font = Font(bold=True, size=12)
    datos = [("Sistema:", proy.sistema), ("Fecha de elaboración:", proy.fecha),
             ("Responsable elaboración:", proy.responsable), ("Flujo (KBDO):", proy.flujo_kbd),
             ("Flujo (bph):", proy.flujo_kbd * KBD_A_BPH), ("Producto:", proy.producto),
             ("Densidad (kg/m3):", proy.densidad), ("Viscosidad (cSt):", proy.viscosidad),
             ("INY 1 (bph):", proy.iny1_bph), ("INY 2 (bph):", proy.iny2_bph)]
    for i, (a, b) in enumerate(datos, 3):
        ws.cell(i, 1, a).font = Font(bold=True)
        ws.cell(i, 2, b)
    f0 = 15
    encabezado(ws, f0, ["Inicio Segmento", "Ubicación (km)", "Presión recibo (psig)", "Presión descarga (psig)",
                        "Flujo (Kbdo)", "Rotura de tubo", "Lleno de línea (kbls)", "Tramo",
                        "Presion Int remanente psig (Descuelgue)", "INYECCIÓN", "%DR", "FR",
                        "Descuelgues", "Re mín", "Régimen"])
    kms = km_estaciones(proy)
    for i, e in enumerate(proy.estaciones):
        r = res[i] if i < len(res) else None
        prec = res[i - 1].presion_recibo if 0 < i <= len(res) else None
        fila = [e.nombre, kms[i], prec]
        if r:
            fila += [r.presion_descarga, r.flujo_kbd, r.roturas, r.lleno_kbls, r.nombre,
                     r.presion_remanente, e.inyeccion, e.dr, e.fr, r.descuelgues, r.re_min, r.regimen]
        for j, v in enumerate(fila, 1):
            ws.cell(f0 + 1 + i, j, v)
    ft = f0 + 2 + len(proy.estaciones)
    ws.cell(ft, 1, "Total (km)").font = Font(bold=True)
    ws.cell(ft, 2, kms[-1])
    ws.cell(ft, 6, "Line fill").font = Font(bold=True)
    ws.cell(ft, 7, sum(r.lleno_kbls for r in res))

    # Interpolación
    fi = ft + 3
    ws.cell(fi - 1, 1, "Interpolación de puntos requeridos").font = Font(bold=True)
    encabezado(ws, fi, ["dist.del reg. [km]", "dist. tramo [km]", "altura [m]", "Presión geodésica (psig)",
                        "MAOP (psig)", "Presión int (psig)", "MAOP remanente (psig)", "Presión int+geod (psig)",
                        "Velocidad [m/s]", "Nre", "Tramo", "Inicio tramo", "Fin tramo", "# Fila"])
    for k, (km, d) in enumerate(zip(proy.km_interpolar, interpolar_lista(proy, res, proy.km_interpolar)), 1):
        if d is None:
            ws.cell(fi + k, 1, km)
            ws.cell(fi + k, 2, "Fuera de rango")
            continue
        for j, key in enumerate(["km", "dist_tramo", "altura", "p_geod", "maop", "p_int", "maop_rem",
                                 "p_int_geod", "velocidad", "nre", "tramo", "inicio", "fin", "fila"], 1):
            ws.cell(fi + k, j, d[key])

    # Hojas de tramo
    cab = ["Punto/ubicación", "dist.del reg. [km]", "altura [m]", "esp. [in]", "diám. ext. [in]", "SMYS [PSI]",
           "Presión geodésica (psig)", "Velocidad [m/s]", "Nre", "Pérdidas de fricción con %DR (psi/km)",
           "Presión int+geod (psig)", "Presión int (psig)", "MAOP (psig)", "MAOP+geod (psig)",
           "Descuelgue", "Rotura tubo", "Lleno de línea", "Rating de Válvula (psi)"]
    for tr, r in zip(proy.tramos, res):
        wt = wb.create_sheet(f"Tramo {tr.nombre}"[:31])
        wt["A1"] = f"Caudal {r.flujo_kbd:.4f} KBDO | %DR {r.dr:.4f} | FR {r.fr:.4f} | Rugosidad {tr.rugosidad_in} in" \
                   f" | Eficiencia {tr.eficiencia}% | P.descarga {r.presion_descarga:.2f} | P.recibo {r.presion_recibo:.2f}"
        encabezado(wt, 2, cab)
        for k, pt in enumerate(tr.puntos):
            vals = [pt.punto, pt.km, pt.altura, pt.espesor, pt.diam_ext, pt.smys, r.J[k], r.K[k], r.L[k],
                    r.M[k], r.N[k], r.O[k], r.P[k], r.Q[k], r.R[k], r.S[k], r.T[k], pt.rating or None]
            for j, v in enumerate(vals, 1):
                c = wt.cell(k + 3, j, v)
                if r.S[k]:
                    c.fill = rojo
                elif r.R[k]:
                    c.fill = amar
        wt.freeze_panes = "B3"

    # Red: nodos de derivación y tramos de cada ramal
    red = getattr(res, "red", None) or {}
    if red:
        wr = wb.create_sheet("Red - ramales")
        encabezado(wr, 1, ["Ramal", "Sale de", "Nodo", "Modo", "% del flujo", "Flujo ramal (KBDO)",
                           "Flujo disponible en el nodo (KBDO)", "Presión en el nodo (psig)",
                           "P. que pide la línea (psig)", "P. que pide el ramal (psig)",
                           "Estrangula línea (psi)", "Estrangula ramal (psi)", "Nota"])
        for i, (nom, inf) in enumerate(red.items(), 2):
            ram = inf["ramal"]
            padre = cadena_de(proy, ram.padre)
            nodo = padre.estaciones[ram.nodo].nombre if padre is not None and ram.nodo < len(padre.estaciones) else ""
            vals = [nom, ram.padre or "Línea principal", nodo, ram.modo, inf["pct_calculado"], inf["flujo"],
                    inf["disponible_nodo"], inf["p_nodo"], inf["d_padre"], inf["d_ramal"], inf["estrangula_padre"],
                    inf["estrangula_ramal"], inf.get("nota", "")]
            for j, v in enumerate(vals, 1):
                wr.cell(i, j, v)
            f = len(red) + 4 + (i - 2) * (len(ram.estaciones) + 3)
            wr.cell(f, 1, f"Ramal {nom}").font = Font(bold=True)
            encabezado(wr, f + 1, ["Estación", "Ubicación (km)", "Presión recibo (psig)", "Presión descarga (psig)",
                                   "Flujo (Kbdo)", "Rotura de tubo", "Lleno de línea (kbls)", "Tramo"])
            pr = proyecto_ramal(proy, ram, inf["flujo"])
            kr = km_estaciones(pr)
            rr = inf["res"]
            for k, e in enumerate(ram.estaciones):
                r = rr[k] if k < len(rr) else None
                fila = [e.nombre, kr[k] if k < len(kr) else None, rr[k - 1].presion_recibo if 0 < k <= len(rr) else None]
                if r:
                    fila += [r.presion_descarga, r.flujo_kbd, r.roturas, r.lleno_kbls, r.nombre]
                for j, v in enumerate(fila, 1):
                    wr.cell(f + 2 + k, j, v)
            for tr, r in zip(ram.tramos, rr):
                if not tr.puntos:
                    continue
                wt = wb.create_sheet(f"{nom[:18]} {tr.nombre}"[:31])
                wt["A1"] = (f"Ramal {nom} | Caudal {r.flujo_kbd:.4f} KBDO | %DR {r.dr:.4f} | "
                            f"P.descarga {r.presion_descarga:.2f} | P.recibo {r.presion_recibo:.2f}")
                encabezado(wt, 2, cab)
                for k, pt in enumerate(tr.puntos):
                    vals = [pt.punto, pt.km, pt.altura, pt.espesor, pt.diam_ext, pt.smys, r.J[k], r.K[k], r.L[k],
                            r.M[k], r.N[k], r.O[k], r.P[k], r.Q[k], r.R[k], r.S[k], r.T[k], pt.rating or None]
                    for j, v in enumerate(vals, 1):
                        c = wt.cell(k + 3, j, v)
                        if r.S[k]:
                            c.fill = rojo
                        elif r.R[k]:
                            c.fill = amar
                wt.freeze_panes = "B3"

    if curva:
        wc = wb.create_sheet("Curva del sistema")
        nombres = curva_nombres or [e.nombre for e in proy.estaciones[:len(proy.tramos)]]
        if curva_titulo:
            wc.insert_rows(1)
        encabezado(wc, 2 if curva_titulo else 1, ["KBDO"] + [f"P. desc. {n} (psig)" for n in nombres])
        if curva_titulo:
            wc.cell(1, 1, curva_titulo).font = Font(bold=True)
        for i, (q, ps) in enumerate(curva, 3 if curva_titulo else 2):
            wc.cell(i, 1, q)
            for j, v in enumerate(ps, 2):
                wc.cell(i, j, v)

    # Gráfico del perfil como imagen
    try:
        from openpyxl.drawing.image import Image as XLImage
        from matplotlib.figure import Figure      # opcional: sin matplotlib no se pone la imagen
        fig = Figure(figsize=(12, 5), dpi=100)
        ax = fig.add_subplot(111)
        for i, (tr, r) in enumerate(zip(proy.tramos, res)):
            km = [p.km for p in tr.puntos]
            ax.plot(km, r.J, color="#8c6d31", label="Presión geodésica" if i == 0 else None)
            ax.plot(km, r.N, color="#1f77b4", label="Presión int+geod" if i == 0 else None)
            ax.plot(km, r.Q, color="#d62728", ls="--", label="MAOP+geod" if i == 0 else None)
        ax.set_xlabel("Distancia (km)")
        ax.set_ylabel("Presión (psi)")
        ax.grid(True, alpha=0.3)
        ax.legend()
        fig.tight_layout()
        buf = io.BytesIO()
        fig.savefig(buf, format="png")
        buf.seek(0)
        img = XLImage(buf)
        ws.add_image(img, "Q3")
    except Exception:  # noqa: BLE001
        pass
    wb.save(ruta)
